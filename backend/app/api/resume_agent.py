"""简历 Agent 统一路由模块。

路由挂载在单个 APIRouter 上，无 /api 前缀；由 main.py 统一加 /agent 前缀。
提示词、工具定义、面试题库、对话逻辑已拆分到 agent/ 子包。
"""
from __future__ import annotations

import io
import json
import os
import re
import random
import asyncio
import base64
import socket
import time
import urllib.request
import urllib.error
import urllib.parse
from collections import Counter, defaultdict
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from fastapi import APIRouter, Depends, HTTPException, File, Query, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.responses import StreamingResponse, HTMLResponse, FileResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from docx import Document
from docx.shared import Pt, Cm
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from ..core.config import DATA_DIR, AVATARS_DIR
from ..services.resume_db import get_db
from ..schemas import resume_models as models
from ..schemas import resume_schemas as schemas
from ..services.resume_provider import (
    stream_response, stream_with_tools, complete_response, load_settings, get_api_key,
    save_settings, test_connection, PROVIDERS,
    add_custom_model, delete_custom_model, update_custom_model,
    get_custom_model, CUSTOM_PROVIDER_KEY, sse_with_heartbeat,
)
from ..services.provider_settings import is_cloud_provider, get_cloud_template_id, DEBUG_LOG_ENABLED
from ..services.resume_parser import parse_resume_pdf_to_background, stream_resume_pdf_to_background
from ..services.resume_html import (
    normalize_structured, structured_to_text, ACCENTS,
)
from ..services.resume_workspace import (
    resume_file_path, ensure_resume_file, read_resume_file, write_resume_file,
    delete_resume_dir, get_workspace_manager,
)
from .agent.prompts import (
    SYSTEM_PROMPT, SYSTEM_PROMPT_AGENT, INTERVIEW_SIM_SYSTEM, INTERVIEW_REPORT_SYSTEM,
    EVAL_SYSTEM_PROMPT, SYSTEM_GENERATE_RESUME, SYSTEM_PARSE_RESUME, SYSTEM_RESTYLE_RESUME,
)
from .agent.tools import RESUME_TOOLS, execute_tool as _execute_tool, _list_sections
from .agent.interview import (
    BankQuestion, load_question_bank, sample_questionnaire, questionnaire_to_markdown,
    generate_job_question_dicts, parse_llm_question_list, QUESTION_CATEGORIES,
    _strip_code_fence,
)
from .agent.chat import generate as _generate, stream_sim as _stream_sim

router = APIRouter(tags=["resume-agent"])


# ── 数据文件路径 ──────────────────────────────────────────────
_LEGACY_DATA_DIR = Path(__file__).resolve().parents[1] / "data"
_legacy_models = _LEGACY_DATA_DIR / "competency_models.json"
_current_models = DATA_DIR / "competency_models.json"
if _legacy_models.exists() and not _current_models.exists():
    import shutil
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(_legacy_models, _current_models)


class GenerateResumeRequest(BaseModel):
    job_id: int
    background: str
    style: Optional[str] = "classic"
    accent: Optional[str] = "#2563eb"


RESTYLE_STYLES = ("classic", "modern", "minimal", "creative", "table", "timeline")


class RestyleResumeRequest(BaseModel):
    style: str = "classic"
    accent: str = "#2563eb"
    description: Optional[str] = ""
    avatar_base64: Optional[str] = None
    reference_image_base64: Optional[str] = None


# ══════════════════════════════════════════════════════════════
# 通用辅助函数
# ══════════════════════════════════════════════════════════════

def _build_job_content(db_job) -> str:
    """拼接岗位完整信息供 Agent 使用。"""
    return db_job.content or ""


def _build_resume_title_with_angle(db_job: models.Job, title: Optional[str], angle: str) -> str:
    base = (title or '').strip()
    if not base:
        base = db_job.title or '岗位'
    if angle and angle not in base:
        return f"{base} · {angle}"
    return base


# ══════════════════════════════════════════════════════════════
# 面试模拟辅助
# ══════════════════════════════════════════════════════════════


def _db_job_questions_as_bank(
    db: Session,
    job_id: int,
    resume_id: Optional[int],
    background_profile_id: Optional[int],
) -> List[BankQuestion]:
    q = db.query(models.JobInterviewQuestion).filter(models.JobInterviewQuestion.job_id == job_id)
    if resume_id is not None:
        q = q.filter(models.JobInterviewQuestion.resume_id == resume_id)
    if background_profile_id is not None:
        q = q.filter(models.JobInterviewQuestion.background_profile_id == background_profile_id)
    rows = q.order_by(models.JobInterviewQuestion.id.asc()).all()
    return [
        BankQuestion(
            id=f"jobq-{r.id}",
            category=r.category,
            text=r.text,
            tags=[],
        )
        for r in rows
    ]


def _format_transcript(messages: list) -> str:
    lines = []
    for m in messages:
        role = m.get("role", "")
        content = (m.get("content") or "").strip()
        if not content:
            continue
        if role == "user":
            lines.append(f"【候选人】\n{content}\n")
        elif role == "assistant":
            lines.append(f"【面试官】\n{content}\n")
    return "\n---\n\n".join(lines)


# ══════════════════════════════════════════════════════════════
# 导出辅助
# ══════════════════════════════════════════════════════════════

def _clean_content(content: str) -> str:
    content = re.sub(r"===RESUME_START===\n?", "", content)
    return re.sub(r"===RESUME_END===\n?", "", content)


def _safe_filename(title: Optional[str], suffix: str, fallback: str) -> str:
    if title:
        safe_title = re.sub(r'[<>:"/\\|?*]', "", title).strip()[:50]
        if safe_title:
            return f"{safe_title}.{suffix}"
    return fallback


def _content_disposition(filename: str) -> str:
    # HTTP 头仅支持 latin-1；非 ASCII 文件名需走 RFC 5987 的 filename* 编码
    quoted = urllib.parse.quote(filename)
    return f"attachment; filename*=UTF-8''{quoted}"


def _strip_md_inline(text: str) -> str:
    """Remove markdown bold/italic/code from text for Word export."""
    s = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    s = re.sub(r"\*(.*?)\*", r"\1", s)
    s = re.sub(r"_(.*?)_", r"\1", s)
    s = re.sub(r"__(.*?)__", r"\1", s)
    s = re.sub(r"`(.*?)`", r"\1", s)
    return s


def md_to_pdf_content(md_text: str, font_size: int = 10):
    """font_size: 9, 10, 11, 12 (body text base)."""
    font_path_regular = None
    font_candidates = [
        os.path.expandvars(r"%WINDIR%\Fonts\msyh.ttc"),
        os.path.expandvars(r"%WINDIR%\Fonts\msyhbd.ttc"),
        os.path.expandvars(r"%WINDIR%\Fonts\simsun.ttc"),
        "C:\\Windows\\Fonts\\msyh.ttc",
        "C:\\Windows\\Fonts\\simsun.ttc",
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    ]
    for fp in font_candidates:
        if not fp:
            continue
        normalized = os.path.normpath(os.path.expanduser(os.path.expandvars(fp)))
        if os.path.exists(normalized):
            font_path_regular = normalized
            break

    if font_path_regular:
        try:
            pdfmetrics.registerFont(TTFont("ChineseFont", font_path_regular))
            base_font = "ChineseFont"
        except Exception:
            base_font = "Helvetica"
    else:
        base_font = "Helvetica"

    getSampleStyleSheet()
    h1_size = max(14, font_size + 8)
    h2_size = max(11, font_size + 3)
    h3_size = max(10, font_size + 1)
    style_h1 = ParagraphStyle("H1", fontName=base_font, fontSize=h1_size, spaceAfter=6, spaceBefore=10, textColor=colors.HexColor("#1a1a2e"), leading=h1_size + 4)
    style_h2 = ParagraphStyle("H2", fontName=base_font, fontSize=h2_size, spaceAfter=4, spaceBefore=8, textColor=colors.HexColor("#16213e"), leading=h2_size + 4)
    style_h3 = ParagraphStyle("H3", fontName=base_font, fontSize=h3_size, spaceAfter=3, spaceBefore=5, textColor=colors.HexColor("#0f3460"), leading=h3_size + 3)
    style_body = ParagraphStyle("Body", fontName=base_font, fontSize=font_size, spaceAfter=2, spaceBefore=1, leading=font_size + 4, textColor=colors.HexColor("#333333"))
    style_bullet = ParagraphStyle("Bullet", fontName=base_font, fontSize=font_size, spaceAfter=2, spaceBefore=1, leading=font_size + 4, leftIndent=15, textColor=colors.HexColor("#333333"))

    story = []
    for line in md_text.split("\n"):
        stripped = re.sub(r"\*\*(.*?)\*\*", r"\1", line.rstrip())
        stripped = re.sub(r"\*(.*?)\*", r"\1", stripped)
        stripped = re.sub(r"`(.*?)`", r"\1", stripped)

        if line.startswith("# "):
            story.append(Paragraph(stripped[2:].strip(), style_h1))
        elif line.startswith("## "):
            story.append(Spacer(1, 0.1 * cm))
            story.append(Paragraph(stripped[3:].strip(), style_h2))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cccccc")))
        elif line.startswith("### "):
            story.append(Paragraph(stripped[4:].strip(), style_h3))
        elif line.startswith("- ") or line.startswith("* "):
            text = stripped[2:].strip()
            if text:
                story.append(Paragraph(f"• {text}", style_bullet))
        elif line.strip() in ["---", "***"]:
            story.append(HRFlowable(width="100%", thickness=0.3, color=colors.HexColor("#dddddd")))
        elif not line.strip():
            story.append(Spacer(1, 0.2 * cm))
        elif stripped.strip():
            story.append(Paragraph(stripped, style_body))

    return story


def _add_paragraph_with_font(doc: Document, text: str, style_name: Optional[str] = None, font_pt: int = 11):
    p = doc.add_paragraph(text, style=style_name)
    for run in p.runs:
        run.font.size = Pt(font_pt)
    return p


# ══════════════════════════════════════════════════════════════
# 文件解析辅助（uploads）
# ══════════════════════════════════════════════════════════════

def _extract_from_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts).strip()


def _extract_from_docx(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _docx_to_pdf_bytes(content: bytes) -> Optional[bytes]:
    """通过 LibreOffice 将 docx 转为 PDF 字节。LibreOffice 不可用时返回 None。"""
    import subprocess, tempfile, os, shutil
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        for p in [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]:
            if os.path.exists(p):
                soffice = p
                break
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "input.docx")
        with open(docx_path, "wb") as f:
            f.write(content)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmpdir, docx_path],
                capture_output=True, timeout=60,
            )
        except Exception:
            return None
        pdf_path = os.path.join(tmpdir, "input.pdf")
        if not os.path.exists(pdf_path):
            return None
        with open(pdf_path, "rb") as f:
            return f.read()


async def _ai_vl_extract_from_pdf(pdf_bytes: bytes) -> str:
    """使用多模态 AI 从 PDF 页面图片中提取文本（与 PDF VL 路径一致）。"""
    from ..services.resume_parser import _pdf_pages_png_base64
    from ..services.resume_provider import is_current_model_multimodal

    settings = load_settings()
    if not is_current_model_multimodal(settings):
        raise HTTPException(
            status_code=400,
            detail="需要多模态模型识别。请在「模型管理」中选择支持多模态的模型。",
        )
    images_b64 = _pdf_pages_png_base64(pdf_bytes, 5)
    if not images_b64:
        return ""
    parts = []
    for b64 in images_b64:
        parts.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}})
    parts.append({
        "type": "text",
        "text": "请识别以上简历页面中的全部文字，按原始结构完整输出。保留所有时间、岗位、项目描述等关键信息。",
    })
    out = await complete_response(
        "你是简历文字识别助手。请完整准确地识别用户提供的简历图片中的所有文字，保持原始结构，不编造、不遗漏。",
        [{"role": "user", "content": parts}],
    )
    return out.strip()


def _ocr_from_image(content: bytes) -> str:
    """图片：交给多模态 LLM 完成文本提取（不使用本地 OCR 库）。"""
    from ..services.resume_provider import complete_response_sync, is_current_model_multimodal
    if not is_current_model_multimodal(load_settings()):
        raise HTTPException(
            status_code=400,
            detail="图片识别需要多模态模型。请在「模型管理」中选择支持多模态的模型。",
        )
    import base64
    b64 = base64.b64encode(content).decode("ascii")
    from PIL import Image
    image = Image.open(io.BytesIO(content))
    fmt = (image.format or "PNG").upper()
    mime = f"image/{fmt.lower()}"
    messages = [
        {
            "role": "user",
            "content": [
                {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                {"type": "text", "text": "请识别这张简历图片中的全部文字，按原始结构完整输出。保留所有时间、岗位、项目描述等关键信息，不要遗漏。"},
            ],
        }
    ]
    try:
        out = complete_response_sync(
            "你是简历文字识别助手。请完整准确地识别用户提供的简历图片中的所有文字，保持原始结构，不编造、不遗漏。",
            messages,
            max_tokens=8192,
            timeout=300.0,
        )
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"多模态模型识别图片失败: {str(e)[:300]}") from e
    if not out.strip():
        raise HTTPException(status_code=502, detail="模型未识别到文字，请检查图片清晰度")
    return out.strip()


def _extract_from_plain_text(content: bytes) -> str:
    """纯文本 / Markdown：按 BOM → UTF-8 → 国标中文编码依次尝试解码。"""
    if not content:
        return ""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return content.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace").strip()


def extract_text_from_file(filename: str, content: bytes):
    lower = filename.lower()
    if lower.endswith((".txt", ".md", ".markdown")):
        return _extract_from_plain_text(content), "text_plain"
    if lower.endswith(".pdf"):
        return _extract_from_pdf(content), "pdf"
    if lower.endswith(".doc") or lower.endswith(".docx"):
        return _extract_from_docx(content), "word"
    if lower.endswith((".png", ".jpg", ".jpeg", ".bmp", ".webp")):
        return _ocr_from_image(content), "ocr"
    raise HTTPException(status_code=400, detail="仅支持 PDF/Word/图片/纯文本(.txt、.md) 文件")


def _parse_job_fields(text: str):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    title = ""
    company = None

    for line in lines[:8]:
        if not title and any(k in line for k in ["工程师", "经理", "专家", "负责人", "Developer", "Manager"]):
            title = re.sub(r"^(职位|岗位)[:：]\s*", "", line)
            continue
        if not company and any(k in line for k in ["公司", "企业", "Company"]):
            company = re.sub(r"^(公司|企业|Company)[:：]\s*", "", line)

    if not title and lines:
        title = lines[0][:80]

    return JobParseResponse(
        title=title or "未识别岗位名称",
        company=company,
        content=text.strip(),
        source="auto_parse",
    )


# ══════════════════════════════════════════════════════════════
# 评分卡辅助（evaluation）
# ══════════════════════════════════════════════════════════════

_COMPETENCY_MODELS: Dict[str, List[str]] = {}


def _load_competency_models() -> Dict[str, List[str]]:
    global _COMPETENCY_MODELS
    if _COMPETENCY_MODELS:
        return _COMPETENCY_MODELS
    fp = DATA_DIR / "competency_models.json"
    try:
        data = json.loads(fp.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            _COMPETENCY_MODELS = {
                str(k): [str(x) for x in v if str(x).strip()]
                for k, v in data.items()
                if isinstance(v, list)
            }
    except Exception:
        _COMPETENCY_MODELS = {}
    return _COMPETENCY_MODELS


def _pick_competency_set(job_title: str) -> List[str]:
    comp_models = _load_competency_models()
    title = (job_title or "").lower()
    if any(k in title for k in ["产品", "product", "pm"]):
        return comp_models.get("product", comp_models.get("default", []))
    if any(k in title for k in ["开发", "工程", "engineer", "后端", "前端"]):
        return comp_models.get("engineering", comp_models.get("default", []))
    if any(k in title for k in ["运营", "operation"]):
        return comp_models.get("operations", comp_models.get("default", []))
    return comp_models.get("default", [])


def _pick_competency_set_by_profile(profile: str, job_title: str) -> List[str]:
    comp_models = _load_competency_models()
    p = (profile or "").strip()
    if p and p in comp_models:
        return comp_models[p]
    return _pick_competency_set(job_title)


def _normalize_scorecard(data: Dict[str, Any]) -> schemas.EvaluationScorecardResponse:
    _SCORE_KEY_MAP = {
        "overall_score": ["overall_score", "total_score", "综合评分", "总分", "score", "总评分"],
        "overall_summary": ["overall_summary", "total_summary", "总结", "综合评价", "summary", "总体评价"],
        "items": ["items", "item_list", "评分项", "competencies", "categories", "dimensions"],
        "needs_verification": ["needs_verification", "待验证", "待核实", "verification_needed", "gaps"],
    }
    _ITEM_SCORE_MAP = ["score", "评分", "rating", "分数", "points"]
    _ITEM_COMPETENCY_MAP = ["competency", "能力项", "能力", "competency_name", "dimension", "name", "title"]

    def _find(data: dict, keys: list, default=None):
        for k in keys:
            if k in data:
                return data[k]
        return default

    overall_score = int(max(0, min(100, int(_find(data, _SCORE_KEY_MAP["overall_score"], 0)))))
    overall_summary = str(_find(data, _SCORE_KEY_MAP["overall_summary"], "")).strip()
    needs_verification = [str(x).strip() for x in (_find(data, _SCORE_KEY_MAP["needs_verification"], []) or []) if str(x).strip()]
    items_raw = _find(data, _SCORE_KEY_MAP["items"], []) or []
    if not isinstance(items_raw, list):
        items_raw = []
    items: List[schemas.ScoreCardItem] = []
    for it in items_raw:
        if not isinstance(it, dict):
            continue
        ev_list = it.get("evidence", []) if isinstance(it.get("evidence", []), list) else []
        evidence: List[schemas.EvidenceRef] = []
        for ev in ev_list:
            if not isinstance(ev, dict):
                continue
            quote = str(ev.get("quote", "")).strip()
            if not quote:
                continue
            evidence.append(
                schemas.EvidenceRef(
                    source=str(ev.get("source", "对话")).strip() or "对话",
                    quote=quote,
                    why=str(ev.get("why", "")).strip() or None,
                )
            )
        score_val = _find(it, _ITEM_SCORE_MAP, 0)
        items.append(
            schemas.ScoreCardItem(
                competency=str(_find(it, _ITEM_COMPETENCY_MAP, "综合能力")).strip() or "综合能力",
                score=int(max(0, min(100, int(score_val)))),
                confidence=str(it.get("confidence", "中")).strip() or "中",
                summary=str(it.get("summary", "")).strip(),
                evidence=evidence,
                gap=str(it.get("gap", "")).strip() or None,
                suggestion=str(it.get("suggestion", "")).strip() or None,
            )
        )
    return schemas.EvaluationScorecardResponse(
        overall_score=overall_score,
        overall_summary=overall_summary,
        items=items,
        needs_verification=needs_verification,
    )


# ══════════════════════════════════════════════════════════════
# Pydantic 模型（settings / uploads）
# ══════════════════════════════════════════════════════════════

class SettingsTest(BaseModel):
    provider: str
    model: str
    api_key: Optional[str] = None
    base_url: Optional[str] = None


class SettingsUpdate(BaseModel):
    provider: str
    model: str
    api_keys: Optional[Dict[str, str]] = None
    custom_base_url: Optional[str] = None
    custom_api_key: Optional[str] = None
    custom_model: Optional[str] = None


class CustomModelCreate(BaseModel):
    """添加自定义模型请求体，所有字段对齐前端 AddModelModal.tsx。"""
    api_format: str = "openai_chat"
    base_url: str
    full_url: bool = False
    model_id: str
    multimodal: bool = False
    api_key: str
    # 高级配置
    model_series: str = "default"
    display_name: str = ""
    context_input: int = 184000
    context_output: int = 16000
    tool_call_rounds: int = 0


class CustomModelTest(BaseModel):
    """连通性测试请求体（前端通过后端代理，避免浏览器 CORS 限制）。
    api_key 为空或 '••••••' 时，若提供了 existing_model_id，则用后端已存储的密钥测试。
    """
    base_url: str
    full_url: bool = False
    model_id: str
    api_key: str = ""
    existing_model_id: Optional[str] = None


class CustomModelUpdate(BaseModel):
    """更新自定义模型请求体，所有字段可选。api_key 留空或 '••••••' 表示保留原值。"""
    api_format: Optional[str] = None
    base_url: Optional[str] = None
    full_url: Optional[bool] = None
    model_id: Optional[str] = None
    multimodal: Optional[bool] = None
    api_key: Optional[str] = None
    model_series: Optional[str] = None
    display_name: Optional[str] = None
    context_input: Optional[int] = None
    context_output: Optional[int] = None
    tool_call_rounds: Optional[int] = None


class ParseResumeBackgroundResponse(BaseModel):
    filename: str
    text: str
    parser: str
    warning: Optional[str] = None


class JobParseResponse(BaseModel):
    title: str
    company: Optional[str] = None
    content: str
    source: str


# ══════════════════════════════════════════════════════════════
# 1. jobs 路由
# ══════════════════════════════════════════════════════════════

@router.get("/jobs/competency-profiles")
def list_competency_profiles():
    fp = DATA_DIR / "competency_models.json"
    try:
        data = json.loads(fp.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            return {"profiles": ["default"]}
        return {"profiles": list(data.keys())}
    except Exception:
        return {"profiles": ["default"]}


@router.get("/jobs", response_model=List[schemas.JobResponse])
def list_jobs(db: Session = Depends(get_db)):
    return db.query(models.Job).order_by(models.Job.created_at.desc()).all()


@router.post("/jobs", response_model=schemas.JobResponse)
def create_job(job: schemas.JobCreate, db: Session = Depends(get_db)):
    db_job = models.Job(**job.model_dump())
    db.add(db_job)
    db.commit()
    db.refresh(db_job)
    return db_job


class JobBatchCreate(BaseModel):
    jobs: List[schemas.JobCreate]


@router.post("/jobs/batch", response_model=List[schemas.JobResponse])
def batch_create_jobs(payload: JobBatchCreate, db: Session = Depends(get_db)):
    results = []
    for item in payload.jobs:
        db_job = models.Job(**item.model_dump())
        db.add(db_job)
        results.append(db_job)
    db.commit()
    for r in results:
        db.refresh(r)
    return results


@router.get("/jobs/{job_id}", response_model=schemas.JobResponse)
def get_job(job_id: int, db: Session = Depends(get_db)):
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")
    return db_job


@router.put("/jobs/{job_id}", response_model=schemas.JobResponse)
def update_job(job_id: int, job: schemas.JobUpdate, db: Session = Depends(get_db)):
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")
    update_data = job.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(db_job, field, value)
    db.commit()
    db.refresh(db_job)
    return db_job


@router.delete("/jobs/{job_id}")
def delete_job(job_id: int, db: Session = Depends(get_db)):
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")
    db.delete(db_job)
    db.commit()
    return {"message": "Job deleted"}


# ══════════════════════════════════════════════════════════════
# 2. resumes 路由
# ══════════════════════════════════════════════════════════════

@router.get("/resumes", response_model=List[schemas.ResumeResponse])
def list_resumes(
    job_id: int = None,
    background_profile_id: int = None,
    db: Session = Depends(get_db),
):
    query = db.query(models.Resume)
    if job_id:
        query = query.filter(models.Resume.job_id == job_id)
    if background_profile_id is not None:
        query = query.filter(
            (models.Resume.background_profile_id == background_profile_id)
            | (models.Resume.background_profile_id.is_(None))
        )
    return query.order_by(models.Resume.created_at.desc()).all()


def _normalize_structured_field(raw: Optional[str]) -> Optional[str]:
    """校验并存的结构化 JSON 字符串，非法时返回 None。"""
    if not raw or not raw.strip():
        return None
    try:
        parsed = normalize_structured(json.loads(raw))
    except (json.JSONDecodeError, TypeError, ValueError):
        return None
    return json.dumps(parsed, ensure_ascii=False)


@router.post("/resumes", response_model=schemas.ResumeResponse)
def create_resume(resume: schemas.ResumeCreate, db: Session = Depends(get_db)):
    db_job = db.query(models.Job).filter(models.Job.id == resume.job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    bg_id = resume.background_profile_id
    angle = (resume.angle or '').strip() if resume.angle else None
    if not angle:
        q = db.query(models.Resume).filter(models.Resume.job_id == resume.job_id)
        if bg_id is not None:
            q = q.filter(models.Resume.background_profile_id == bg_id)
        else:
            q = q.filter(models.Resume.background_profile_id.is_(None))
        angle_index = q.count() + 1
        angle = f"角度{angle_index}"

    existing_count = db.query(models.Resume).filter(models.Resume.job_id == resume.job_id)
    if bg_id is not None:
        existing_count = existing_count.filter(models.Resume.background_profile_id == bg_id)
    else:
        existing_count = existing_count.filter(models.Resume.background_profile_id.is_(None))
    existing_count = existing_count.filter(models.Resume.angle == angle).count()

    data = resume.model_dump(exclude={'angle', 'background_profile_id', 'title'})
    data["structured"] = _normalize_structured_field(data.get("structured"))
    content = (data.get("content") or "").strip()
    if not content and data["structured"]:
        content = structured_to_text(json.loads(data["structured"]))
    data["content"] = content

    db_resume = models.Resume(
        **data,
        angle=angle,
        background_profile_id=bg_id,
        title=_build_resume_title_with_angle(db_job, resume.title, angle),
        version=existing_count + 1,
    )
    db.add(db_resume)
    db.commit()
    db.refresh(db_resume)
    return db_resume


@router.get("/resumes/{resume_id}", response_model=schemas.ResumeResponse)
def get_resume(resume_id: int, db: Session = Depends(get_db)):
    db_resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")
    return db_resume


@router.put("/resumes/{resume_id}", response_model=schemas.ResumeResponse)
def update_resume(resume_id: int, resume: schemas.ResumeUpdate, db: Session = Depends(get_db)):
    db_resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")
    update_data = resume.model_dump(exclude_unset=True)
    if "structured" in update_data:
        normalized = _normalize_structured_field(update_data["structured"])
        update_data["structured"] = normalized
        if normalized and "content" not in update_data:
            update_data["content"] = structured_to_text(json.loads(normalized))
    for field, value in update_data.items():
        setattr(db_resume, field, value)
    db.commit()
    db.refresh(db_resume)

    if db_resume.rendered_html:
        write_resume_file(resume_id, db_resume.rendered_html)

    return db_resume


@router.delete("/resumes/{resume_id}")
def delete_resume(resume_id: int, db: Session = Depends(get_db)):
    db_resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")
    db.delete(db_resume)
    db.commit()
    disk_ok = delete_resume_dir(resume_id)
    if not disk_ok:
        raise HTTPException(status_code=500, detail="简历记录已删除但磁盘文件清理失败，请手动清理")
    return {"message": "Resume deleted"}


# 生成阶段可选简历布局：系统提示已不含任何具体布局，每种布局互斥、每次请求只注入其中一种
_GENERATE_LAYOUT_HINTS = {
    "classic": ("""整体布局 — 深色侧边栏双栏：
- 页面容器 .page：width: 210mm; height: 297mm; display: flex; overflow: hidden; background: #fff; box-shadow 预览用
- 左侧深色侧边栏 .sidebar：width: 58mm; min-width: 58mm; height: 100%; background: linear-gradient(160deg, #0d1b2a 0%, #16304d 50%, #1b3a5c 100%); color: #e2e8f0; padding: 7mm 5mm 6mm 5mm; flex-direction: column; position: relative
- 侧边栏右侧加 3px 渐变竖线装饰（::after 伪元素，#00b4d8 → #0077b6）
- 右侧主内容区 .main：flex: 1; height: 100%; padding: 7mm 7mm 6mm 7mm; overflow: hidden

侧边栏内容 — 从上到下：
1. 圆形头像框 .avatar-frame：28mm × 28mm，圆形，border: 3px solid #00b4d8，深色渐变背景，居中，阴影 box-shadow: 0 0 16px rgba(0,180,216,0.25)；有头像时用 <img>，无头像时显示姓名首字（.avatar-placeholder，36px 粗体 #90e0ef）
2. 姓名区 .name-block：姓名 h1（26px 700 #fff，letter-spacing: 3px），下方求职意向 .title（13px #90e0ef，letter-spacing: 1px）
3. 分割线 .sidebar-divider：1.5px 高，背景 linear-gradient(90deg, #00b4d8, transparent)
4. 联系方式 .sb-section：标题 .sb-title（14px 700 #fff，左侧 3px #00b4d8 竖条装饰，右侧英文小标签 .en 如 "CONTACT"）；每条 .contact-item：Bootstrap Icons 图标 + 文字（11.5px #cbd5e1）；在 <head> 引入 Bootstrap Icons CDN：<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/bootstrap-icons@1.11.3/font/bootstrap-icons.min.css">；图标用 <i class="bi bi-xxx"> 写法，颜色 color: #48cae4，字号 font-size: 14px；图标 class 对照：电话 bi-telephone、邮箱 bi-envelope、GitHub bi-github、地址 bi-geo-alt
5. 教育背景 .sb-section：学校名 .edu-school（13px 600 #f1f5f9）、时间 .edu-date（11.5px #94a3b8）、专业 .edu-major（12px #cbd5e1）、补充信息 .edu-meta（11.5px #94a3b8）
6. 专业技能 .sb-section：每类 .skill-group 含 .skill-label（12px 700 #48cae4）+ .skill-val（11.5px #e2e8f0）

主内容区 — 从上到下：
1. 项目经历：大标题 .main-sec-title（17px 700 #0d1b2a，左侧 4px 渐变竖条装饰，右侧英文副标签如 "PROJECTS"），下方 2px solid #e2e8f0 横线；每个项目 .project：左侧 2.5px solid #e2e8f0 竖线；项目头 .project-head：项目名 .project-name（14.5px 700 #0d1b2a）+ 角色 .role（12.5px #0077b6，前空7px） + 日期 .project-date（12.5px #64748b，右对齐）；技术标签 .project-tech：每个 .tag 是 inline-block，背景 #f0f7fb，border 1px solid #c5e3f0，color #0369a1，圆角 3px，padding 1px 7px；要点列表 ul li：13px #334155，左侧 5px #00b4d8 圆点，关键词 <b> 用 #0d1b2a 600
2. 校园经历 .campus-section：大标题同样用 .main-sec-title（右侧英文副标签如 "CAMPUS"），下方复用 .project 结构
3. 自我评价 .eval-section：顶部 2px solid #e2e8f0 分割线，标题 .eval-title 同样带竖条装饰，内容区 .eval-text：13px #334155，line-height: 1.7

配色体系：
- 侧边栏背景：linear-gradient(160deg, #0d1b2a, #16304d, #1b3a5c)；强调色/装饰：#00b4d8（青蓝）用于头像边框、竖条装饰、圆点、标签
- 侧边栏文字层级：标题 #fff，强调 #48cae4 / #90e0ef，正文 #cbd5e1，次要 #94a3b8
- 主区域文字层级：标题 #0d1b2a，正文 #334155，次要 #64748b，分割线 #e2e8f0
- 标签背景：#f0f7fb，标签边框：#c5e3f0，标签文字：#0369a1

字体与间距：全局字号 13px，行高 1.6；姓名 26px 700 letter-spacing: 3px；模块标题：侧边栏 14px 700，主区域 17px 700；项目名 14.5px 700，正文要点 13px；技术标签 11.5px；间距整体紧凑：主区域模块间距 3mm、条目/要点间距 2mm；侧边栏各信息组间距 3.5mm、组内条目间距 1.5mm"""),
    "modern": ("现代极简单栏布局，无卡片无阴影，靠间距和细横线分割各个模块；"
               "页面居中，固定 A4 宽度（内容区约 170mm 水平居中，上下留白约 10mm）；"
               "顶部个人信息区采用左右两栏排布（display: flex，垂直居中对齐）："
               "左栏放文字信息——姓名（24px 700 #111，左对齐）、下方依次一行小字联系方式（用 | 分隔电话/邮箱/地址等）、求职意向、教育背景；"
               "右栏放圆形头像（28mm 正圆，object-fit: cover），靠页面右侧；"
               "无头像时右栏显示姓名首字圆形占位；求职意向和教育背景属于个人信息区内容，紧跟联系方式下方列在左栏内，不单独拆成页面更下方的独立大模块；"
               "下方剩余模块顺序：项目经历、实习经历、校园经历、荣誉证书、自我评价，不得包含其他板块，无对应内容的模块可省略但顺序不得打乱；"
               "候选人有校园经历/学生工作/竞赛时，作为「校园经历」板块紧跟项目经历独立成块（或并入实习经历板块），必须保留不得省略；"
               "视觉：黑白灰配色（标题 #111、正文 #333、次要 #777），细灰横线 1px #ddd 作为区块分隔；"
               "模块标题加粗、字号略大（15px 700 #111），正文 13px、行高 1.7 舒适；"
               "无任何装饰元素：无图标、无色块、无渐变、无主题色竖条；--accent 变量仍需定义，但页面视觉不依赖主题色"),
    "minimal": ("单栏居中 max-width: 160mm，大量留白，padding: 20mm；姓名居中（20px），下方仅一条宽 40mm 细主题色短线；"
                "无图标、无背景色块、无侧边栏、无分割线；模块标题用主题色文字 + 1px 浅灰线；正文 12px，行距 1.8；"
                "技能用逗号分隔文本，联系方式一行用 | 分隔"),
    "creative": ("顶部全幅主题色色带（高约 80mm）：白色大号姓名（28px）+ 求职意向（16px），头像嵌入色带右侧圆形白边框；"
                 "下方双栏：左 40% 浅灰背景 #f5f5f5 放技能/荣誉/教育，右 60% 白色放工作/项目/校园经历/自我评价；"
                 "模块标题深色小号加粗 + 2px 主题色左竖条；正文 12px，行距 1.6"),
    "table": ("使用表格 <table> 布局简历，正式商务风格，白底黑字，细灰色分割线（1px #ccc），不要花哨渐变、阴影；"
              "表格结构：第一行大标题（姓名 + 求职意向）跨列合并；第二行个人信息一行，单元格分别放电话、邮箱、地址、GitHub；"
              "下面依次是教育经历、项目经历、技能证书、自我评价，每个大标题单独占一行跨列合并，内容放在子表格单元格内；"
              "候选人有校园经历/学生工作/竞赛时，在项目经历之后增加一行「校园经历」大标题行（结构同项目经历），必须保留不得省略；"
              "字体用系统无衬线字体，字号适中（正文约 12.5px），边框细浅不刺眼；"
              "表格仅用于排版，语义尽量整洁，除大标题行跨列外不要多余合并单元格；"
              "未提供的信息用合理中文占位符填充；border-collapse: collapse，@media print 时边框不丢失（设置 -webkit-print-color-adjust: exact 和 print-color-adjust: exact）；"
              "--accent 变量仍需定义，但配色以黑白灰为主，主题色不作为装饰"),
    "timeline": ("时间轴风格单栏布局：页面居中，内容区约 170mm 水平居中；"
                 "顶部大姓名（22px 700）+ 下方一行小字联系方式；"
                 "经历类模块（教育背景、项目经历、实习经历/校园经历）使用时间轴呈现——模块内容区左侧一条 2px 浅灰竖向轴线贯穿，"
                 "每条经历对应轴上一个主题色圆点节点，节点旁标注时间，右侧放经历名称、角色与要点列表；"
                 "技能、证书、自我评价等非时间模块正常分块排版（标题 + 内容），不使用时间轴；"
                 "候选人有校园经历时同样按时间轴呈现并必须保留；"
                 "主题色仅用于时间轴节点圆点与少量强调；白底深灰字，正文 12.5-13px，行距 1.6，无渐变无阴影"),
}


def _append_layout_directive(user_msg: str, style: str, accent: str) -> str:
    """把用户选的简历布局作为本次唯一排版指令追加到生成请求（未知/缺省一律用 classic）。"""
    key = (style or "").strip()
    hint = _GENERATE_LAYOUT_HINTS.get(key) or _GENERATE_LAYOUT_HINTS["classic"]
    return (user_msg +
            f"\n\n---\n\n## 排版布局（本次唯一采用的布局，严格按此执行，不得混入其它布局风格）\n\n{hint}\n\n"
            f"主题色 {accent} 写入 CSS 变量 --accent（若上方布局明确要求黑白灰配色，则 --accent 仅作变量预留、不用于视觉装饰）；"
            f"无论采用哪种布局，系统提示中的【板块内容要求】（包含「校园经历」为必保留独立板块）、【防溢出与打印】【HTML 结构要求】等内容规则均不变，仍必须遵守。")


def _merge_profile_background(db: Session, background: str) -> str:
    """把「背景档案」（个人信息表单）自动注入候选人背景，用于生成简历。"""
    prof = db.query(models.UserBackground).order_by(models.UserBackground.id.asc()).first()
    profile_text = (prof.content or "").strip() if prof else ""
    if not profile_text or profile_text in background:
        return background
    return f"{profile_text}\n\n{background}" if background else profile_text


def _avatar_directive(background: str) -> str:
    """从候选人背景文本中提取「头像：本地绝对路径」，生成头像引用指令。

    头像统一存放在安装目录 data/avatars 下；生成简历时把绝对路径发给模型，
    让其在头像框中用 <img src="绝对路径"> 原样引用。预览（webview file:///
    http 内联）与导出层已有本地路径图片的渲染兜底，模型侧无需转 base64。
    """
    if not background:
        return ""
    m = re.search(r"^头像[：:]\s*(.+?)\s*$", background, flags=re.MULTILINE)
    if not m:
        return ""
    raw = m.group(1).strip().strip('"').strip("'")
    p = Path(raw)
    if not p.is_file():
        return ""
    abs_path = str(p.resolve())
    return (
        f"\n\n---\n\n## 候选人头像\n\n"
        f"候选人头像文件的本地绝对路径：`{abs_path}`\n"
        f"必须在简历头像框中原样引用该路径：`<img src=\"{abs_path}\">`。"
        f"禁止改写为相对路径、base64 或虚构的网络 URL；"
        f"若当前布局没有头像框，则在姓名旁新增圆形头像框放置该图片（正圆、object-fit: cover）。"
    )


# AI 生成简历调试日志目录（与 ConversationDebugLog 同目录：backend/debug_logs）
_GENERATE_DEBUG_LOG_DIR = Path(__file__).resolve().parents[2] / "debug_logs"


def _write_generate_debug_log(system: str, user_msg: str, response: str, error: str = "") -> None:
    """把一次「AI 生成简历」的完整请求/响应落盘为 generate_resume_*.log。

    格式沿用旧日志：===== SYSTEM PROMPT / USER MESSAGE / LLM RESPONSE =====，
    受 DEBUG_LOG_ENABLED 开关控制；写日志失败不影响简历生成主流程。
    """
    if not DEBUG_LOG_ENABLED:
        return
    try:
        from datetime import datetime
        from ..services.provider_core import _resolve_model_info_for_log
        provider, model = _resolve_model_info_for_log()
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]
        content = (
            f"===== META =====\n\n"
            f"时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | provider: {provider} | model: {model}\n\n"
            f"===== SYSTEM PROMPT =====\n\n{system}\n\n"
            f"===== USER MESSAGE =====\n\n{user_msg}\n\n"
            f"===== {'LLM ERROR' if error and not response else 'LLM RESPONSE'} =====\n\n"
            f"{response if response else (f'[生成失败] {error}')}\n"
        )
        _GENERATE_DEBUG_LOG_DIR.mkdir(exist_ok=True)
        (_GENERATE_DEBUG_LOG_DIR / f"generate_resume_{stamp}.log").write_text(content, encoding="utf-8")
    except Exception:
        pass


@router.post("/resumes/generate", response_model=schemas.ResumeResponse)
async def generate_resume(body: GenerateResumeRequest, db: Session = Depends(get_db)):
    """AI 根据候选人背景 + 目标岗位 JD 直接生成完整简历 HTML，写入磁盘。"""
    db_job = db.query(models.Job).filter(models.Job.id == body.job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    jd = _build_job_content(db_job).strip()
    if not jd:
        raise HTTPException(status_code=400, detail="岗位 JD 为空，无法生成简历")

    background = _merge_profile_background(db, (body.background or "").strip())
    user_msg = f"## 目标岗位 JD\n\n{jd}"
    if background:
        user_msg += f"\n\n---\n\n## 候选人背景信息\n\n{background}"
    else:
        user_msg += "\n\n---\n\n## 候选人背景信息\n\n（用户未提供候选人背景，请基于目标岗位 JD 生成一份完整的简历模板，姓名、联系方式等字段用合理的占位信息填充。）"
    accent = body.accent if body.accent in ACCENTS else "#2563eb"
    user_msg += f"\n\n---\n\n请使用主题色 {accent}，生成一份完整简历 HTML。"
    user_msg = _append_layout_directive(user_msg, body.style, accent)
    user_msg += _avatar_directive(background)


    try:
        raw = await complete_response(SYSTEM_GENERATE_RESUME, [{"role": "user", "content": user_msg}])
    except Exception as e:
        _write_generate_debug_log(SYSTEM_GENERATE_RESUME, user_msg, "", error=str(e))
        raise
    _write_generate_debug_log(SYSTEM_GENERATE_RESUME, user_msg, raw)
    html = _strip_code_fence(raw).strip()

    if not html or ("<!DOCTYPE" not in html and "<html" not in html):
        raise HTTPException(status_code=502, detail="模型未生成有效的 HTML 简历，请重试或更换模型")

    q = db.query(models.Resume).filter(models.Resume.job_id == body.job_id)
    angle_index = q.count() + 1
    angle = f"角度{angle_index}"
    existing_count = q.filter(models.Resume.angle == angle).count()
    title = _build_resume_title_with_angle(db_job, None, angle)

    db_resume = models.Resume(
        job_id=body.job_id,
        title=title,
        content=background,
        structured=None,
        rendered_html=html,
        current_template=body.style or "classic",
        angle=angle,
        version=existing_count + 1,
    )
    db.add(db_resume)
    db.commit()
    db.refresh(db_resume)

    write_resume_file(db_resume.id, html)

    return db_resume


@router.post("/resumes/generate/stream")
async def generate_resume_stream(body: GenerateResumeRequest, db: Session = Depends(get_db)):
    """流式版：AI 根据候选人背景 + 目标岗位 JD 逐步生成完整简历 HTML。"""
    db_job = db.query(models.Job).filter(models.Job.id == body.job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    jd = _build_job_content(db_job).strip()
    if not jd:
        raise HTTPException(status_code=400, detail="岗位 JD 为空，无法生成简历")

    background = _merge_profile_background(db, (body.background or "").strip())
    user_msg = f"## 目标岗位 JD\n\n{jd}"
    if background:
        user_msg += f"\n\n---\n\n## 候选人背景信息\n\n{background}"
    else:
        user_msg += "\n\n---\n\n## 候选人背景信息\n\n（用户未提供候选人背景，请基于目标岗位 JD 生成一份完整的简历模板，姓名、联系方式等字段用合理的占位信息填充。）"
    accent = body.accent if body.accent in ACCENTS else "#2563eb"
    user_msg += f"\n\n---\n\n请使用主题色 {accent}，生成一份完整简历 HTML。"
    user_msg = _append_layout_directive(user_msg, body.style, accent)
    user_msg += _avatar_directive(background)


    async def _stream_gen():
        html_parts: list[str] = []
        try:
            async for text in stream_response(SYSTEM_GENERATE_RESUME, [{"role": "user", "content": user_msg}]):
                html_parts.append(text)
                yield f"data: {json.dumps({'type': 'text', 'content': text})}\n\n"
        except Exception as e:
            _write_generate_debug_log(SYSTEM_GENERATE_RESUME, user_msg, "".join(html_parts), error=str(e))
            yield f"data: {json.dumps({'type': 'error', 'content': f'模型输出中断: {str(e)[:300]}'})}\n\n"
            return

        full_html = "".join(html_parts).strip()
        full_html = _strip_code_fence(full_html)
        _write_generate_debug_log(SYSTEM_GENERATE_RESUME, user_msg, full_html)

        if full_html and ("<!DOCTYPE" in full_html or "<html" in full_html):
            q = db.query(models.Resume).filter(models.Resume.job_id == body.job_id)
            angle_index = q.count() + 1
            angle = f"角度{angle_index}"
            existing_count = q.filter(models.Resume.angle == angle).count()
            title = _build_resume_title_with_angle(db_job, None, angle)

            db_resume = models.Resume(
                job_id=body.job_id,
                title=title,
                content=background,
                structured=None,
                rendered_html=full_html,
                current_template=body.style or "classic",
                angle=angle,
                version=existing_count + 1,
            )
            db.add(db_resume)
            db.commit()
            db.refresh(db_resume)

            write_resume_file(db_resume.id, full_html)

            yield f"data: {json.dumps({'type': 'done', 'resume_id': db_resume.id}, ensure_ascii=False)}\n\n"
        else:
            yield f"data: {json.dumps({'type': 'error', 'content': '模型未生成有效的 HTML 简历，请重试或更换模型'})}\n\n"

    return StreamingResponse(
        sse_with_heartbeat(_stream_gen()),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.post("/resumes/{resume_id}/restyle/stream")
async def restyle_resume_stream(resume_id: int, body: RestyleResumeRequest, db: Session = Depends(get_db)):
    """流式重新排版简历：基于已有简历数据，按指定风格重新生成 HTML，更新当前简历。"""
    db_resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")

    resume_content = _extract_resume_text(db_resume)
    resume_id_val = db_resume.id

    style = body.style if body.style in RESTYLE_STYLES else "classic"
    accent = body.accent if body.accent in ACCENTS else "#2563eb"

    avatar_path = None
    if body.avatar_base64:
        import base64
        avatar_dir = resume_dir(resume_id)
        avatar_dir.mkdir(parents=True, exist_ok=True)
        data_uri = body.avatar_base64
        if data_uri.startswith("data:"):
            header, b64data = data_uri.split(",", 1)
            mime = header.split(":")[1].split(";")[0]
            ext = mime.split("/")[-1].replace("jpeg", "jpg")
        else:
            b64data = data_uri
            ext = "png"
            mime = "image/png"
        avatar_file = avatar_dir / f"avatar.{ext}"
        avatar_file.write_bytes(base64.b64decode(b64data))
        avatar_path = str(avatar_file.resolve())

    text_content = f"## 简历原始信息\n\n{resume_content}\n\n---\n\n"
    text_content += f"目标风格：{style}\n主题色：{accent}\n"
    if body.description and body.description.strip():
        text_content += f"用户补充要求：{body.description.strip()}\n"

    has_ref_image = bool(body.reference_image_base64)
    if has_ref_image:
        from ..services.resume_provider import is_current_model_multimodal
        settings = load_settings()
        if not is_current_model_multimodal(settings):
            raise HTTPException(
                status_code=400,
                detail="当前模型不支持多模态，无法分析参考风格图。请切换到支持多模态的模型（如 Claude、GPT-4o 等），或不使用参考图。",
            )
        text_content += "用户已提供参考风格图，请分析其排版风格并融入生成。\n"

    if avatar_path:
        import base64 as _b64
        _ap = Path(avatar_path)
        _avatar_b64 = _b64.b64encode(_ap.read_bytes()).decode()
        _mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp"}.get(_ap.suffix.lower(), "image/png")
        _data_uri = f"data:{_mime};base64,{_avatar_b64}"
        text_content += f"\n用户头像已以 base64 嵌入，请直接使用 <img src=\"{_data_uri}\"> 引用头像。\n"

    text_content += "\n请根据以上信息重新排版生成完整简历 HTML。"

    db_restyle_job = db.query(models.Job).filter(models.Job.id == db_resume.job_id).first()
    restyle_jd = _build_job_content(db_restyle_job) if db_restyle_job else ""

    user_msg_content: Any = text_content
    if has_ref_image and body.reference_image_base64:
        user_msg_content = [
            {"type": "image_url", "image_url": {"url": body.reference_image_base64}},
            {"type": "text", "text": text_content},
        ]

    async def _stream_restyle():
        html_parts: list[str] = []
        try:
            msg = [{"role": "user", "content": user_msg_content}]
            async for text in stream_response(SYSTEM_RESTYLE_RESUME, msg):
                html_parts.append(text)
                yield f"data: {json.dumps({'type': 'text', 'content': text})}\n\n"
        except Exception as e:
            yield f"data: {json.dumps({'type': 'error', 'content': f'模型输出中断: {str(e)[:300]}'})}\n\n"
            return

        full_html = "".join(html_parts).strip()
        full_html = _strip_code_fence(full_html)

        if full_html and ("<!DOCTYPE" in full_html or "<html" in full_html):
            from ..services.resume_db import SessionLocal
            save_db = SessionLocal()
            try:
                save_resume = save_db.query(models.Resume).filter(models.Resume.id == resume_id_val).first()
                if save_resume:
                    save_resume.rendered_html = full_html
                    save_resume.current_template = style
                    save_db.commit()
            finally:
                save_db.close()

            write_resume_file(resume_id_val, full_html)

            yield f"data: {json.dumps({'type': 'done', 'resume_id': resume_id_val}, ensure_ascii=False)}\n\n"
        else:
            yield f"data: {json.dumps({'type': 'error', 'content': '模型未生成有效的 HTML 简历，请重试或更换模型'})}\n\n"

    return StreamingResponse(
        sse_with_heartbeat(_stream_restyle()),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


def _extract_resume_text(db_resume) -> str:
    """从 Resume 数据库记录中提取文本内容，用于重新排版。"""
    if db_resume.structured:
        try:
            return structured_to_text(json.loads(db_resume.structured))
        except Exception:
            pass
    if db_resume.content and db_resume.content.strip():
        return db_resume.content.strip()
    if db_resume.rendered_html:
        from html.parser import HTMLParser
        class _TextExtractor(HTMLParser):
            def __init__(self):
                super().__init__()
                self.parts: list[str] = []
            def handle_data(self, data):
                t = data.strip()
                if t:
                    self.parts.append(t)
        ext = _TextExtractor()
        ext.feed(db_resume.rendered_html)
        return "\n".join(ext.parts)
    return ""


# ══════════════════════════════════════════════════════════════
# 3. chat 路由
# ══════════════════════════════════════════════════════════════

@router.post("/chat/stream")
async def chat_stream(request: schemas.ChatRequest, db: Session = Depends(get_db)):
    db_job = db.query(models.Job).filter(models.Job.id == request.job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    db_resume = db.query(models.Resume).filter(models.Resume.id == request.resume_id).first()
    resume_content = db_resume.content if db_resume else ""
    resume_structured = db_resume.structured if db_resume else None
    resume_rendered_html = db_resume.rendered_html if db_resume else None

    messages = [msg.model_dump() for msg in request.messages]

    return StreamingResponse(
        sse_with_heartbeat(_generate(
            job_content=_build_job_content(db_job),
            resume_content=resume_content,
            messages=messages,
            user_background=_merge_profile_background(db, (request.user_background or "").strip()),
            resume_structured=resume_structured,
            resume_rendered_html=resume_rendered_html,
            resume_id=request.resume_id,
        )),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.websocket("/ws/resume/{resume_id}")
async def resume_ws(ws: WebSocket, resume_id: int):
    """简历文件变更 WebSocket：AI 写入文件后自动推送 file_changed 事件。"""
    await ws.accept()
    mgr = get_workspace_manager()
    mgr.subscribe(resume_id, ws)
    try:
        while True:
            await ws.receive_text()  # 保活，忽略客户端消息
    except WebSocketDisconnect:
        pass
    finally:
        mgr.unsubscribe(resume_id, ws)


@router.get("/resumes/{resume_id}/file")
def get_resume_file(resume_id: int):
    """获取简历磁盘文件的最新内容（前端收到 WebSocket 通知后调用）。"""
    content = read_resume_file(resume_id)
    file_path = str(resume_file_path(resume_id))
    return {"resume_id": resume_id, "html": content, "path": file_path}


@router.post("/chat/rollback")
def rollback_snapshot(request: dict):
    """回滚简历快照到指定回合深度。前端用户点击消息旁的撤销按钮时调用。

    请求体: { "resume_id": int, "target_depth": int }
    - target_depth: 该用户消息发送前已保存的回合快照深度（snapshot_depth）
    - 回合快照是「回合开始前」的文件状态；恢复后丢弃其后所有回合快照
    - 若 target_depth == 0，恢复到对话起点的快照并清空整栈
    """
    from ..services.local_edit import snapshot as snap_mod
    from ..services.resume_workspace import read_resume_file, write_resume_file
    from ..services.resume_db import SessionLocal
    from ..schemas import resume_models as models

    resume_id = request.get("resume_id")
    target_depth = request.get("target_depth", 0)
    if not resume_id:
        raise HTTPException(status_code=400, detail="resume_id is required")

    resume_id = int(resume_id)
    result = snap_mod.rollback_to_round(resume_id, target_depth)

    content = read_resume_file(resume_id)
    write_resume_file(resume_id, content)

    try:
        db = SessionLocal()
        try:
            db_resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
            if db_resume:
                db_resume.rendered_html = content
                db.commit()
        finally:
            db.close()
    except Exception:
        pass

    return {
        "ok": result["ok"],
        "message": result["message"],
        "popped": result["popped"],
        "current_depth": result["current_depth"],
    }


def _inline_local_image_refs(html: str) -> str:
    """将 HTML 中引用的本地绝对路径图片内联为 base64 data URI。

    内置浏览器（webview）通过 http 加载简历 HTML，此时 `<img src="C:\\...">` 中
    的 `C:` 会被当作 URL 协议名而找不到，导致头像/图片无法显示；而直接双击用系统
    浏览器打开（file://）时又能显示。这里在预览输出阶段做统一兜底：
      - 匹配 `<img ... src="本地路径">` 与 CSS `url(本地路径)`
      - Windows 盘符绝对路径（`C:\...` / `C:/...`）或 `file:///C:/...` 视为本地文件
      - 读取并转成 `data:<mime>;base64,...` 原地替换，读不到则保持原样

    仅对 img/src 与 url() 生效，不改写 http/data 等网络资源。
    """
    if not html:
        return html

    _MIME = {
        ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
        ".gif": "image/gif", ".webp": "image/webp", ".svg": "image/svg+xml",
        ".ico": "image/x-icon", ".bmp": "image/bmp",
    }
    _cache: dict[str, str] = {}

    def _local_path(raw: str) -> str | None:
        """把属性值解析为本地绝对路径；非本地路径返回 None。"""
        v = raw.strip()
        if not v or v.startswith(("http://", "https://", "data:", "//", "#")):
            return None
        if v.lower().startswith("file:///"):
            v = v[8:]
        elif v.lower().startswith("file:"):
            v = v[5:]
        m = re.match(r"^([a-zA-Z]):([\\/].*)$", v)
        if not m:
            return None
        # file:/// URL 常对中文等字符做百分号编码，需解码后才是真实文件名
        return urllib.parse.unquote(v)

    def _to_data_uri(path_str: str) -> str:
        """读取本地图片并转 data URI，带进程内缓存。读不到返回原路径。"""
        if path_str in _cache:
            return _cache[path_str]
        try:
            pb = Path(path_str)
            if pb.is_file():
                mime = _MIME.get(pb.suffix.lower(), "image/png")
                data = f"data:{mime};base64," + base64.b64encode(pb.read_bytes()).decode("ascii")
                _cache[path_str] = data
                return data
        except Exception:
            pass
        _cache[path_str] = path_str
        return path_str

    def _replace_local(raw: str) -> str:
        local = _local_path(raw)
        if local is None:
            return raw
        return _to_data_uri(local)

    def _sub_img(m: "re.Match[str]") -> str:
        quote, inner = m.group(1), m.group(2)
        return f'<img src={quote}{_replace_local(inner)}{quote}'

    def _sub_url(m: "re.Match[str]") -> str:
        quote, inner = m.group(1), m.group(2)
        return f'url({quote}{_replace_local(inner)}{quote})'

    # ---- 应用替换 ----
    html = re.sub(r'<img\s+src=(["\'])(.*?)\1', _sub_img, html, flags=re.IGNORECASE)
    html = re.sub(r'url\(\s*(["\'])(.*?)\1\s*\)', _sub_url, html, flags=re.IGNORECASE)
    return html


@router.get("/resumes/{resume_id}/resume.html")
def get_resume_html_file(resume_id: int):
    """输出简历的原始 HTML，供内置浏览器预览（Electron webview）直接加载。

    为了让内置浏览器（http 上下文）与系统浏览器都能正常显示本地绝对路径的图片，
    预览输出前会对 HTML 中引用的本地磁盘图片（Windows 盘符路径 / file:// 前缀）
     统一内联为 base64 data URI，从而不依赖 file:// 协议的读取权限。
    """
    content = read_resume_file(resume_id)
    if not content:
        raise HTTPException(status_code=404, detail="Resume file not found")
    return HTMLResponse(_inline_local_image_refs(content))


@router.get("/chat/current-provider")
def get_current_provider():
    """Return the currently configured provider and model info for the UI."""
    settings = load_settings()
    provider_id = settings.get("provider", "")
    model_id = settings.get("model", "")

    # 自定义模型：从 settings.custom_models 解析显示名
    model_name = model_id
    provider_name = provider_id or "未配置"
    if provider_id.startswith(f"{CUSTOM_PROVIDER_KEY}:"):
        cm = get_custom_model(provider_id, settings)
        if cm:
            model_name = cm.get("display_name") or cm.get("model_id") or model_id
            provider_name = "自定义模型"
    elif provider_id == CUSTOM_PROVIDER_KEY:
        provider_name = "自定义模型"

    return {
        "provider": provider_id,
        "provider_name": provider_name,
        "model": model_id,
        "model_name": model_name,
    }


@router.get("/chat/conversations/{resume_id}")
def get_conversation(resume_id: int, db: Session = Depends(get_db)):
    """Get conversation for a resume. Returns empty messages if none saved."""
    db_conv = (
        db.query(models.Conversation)
        .filter(models.Conversation.resume_id == resume_id)
        .first()
    )
    if not db_conv:
        return {"resume_id": resume_id, "messages": []}
    return {"resume_id": resume_id, "messages": json.loads(db_conv.messages)}


@router.post("/chat/conversations")
def save_conversation(body: schemas.ConversationSave, db: Session = Depends(get_db)):
    resume_id = body.resume_id
    messages = body.messages
    db_conv = (
        db.query(models.Conversation)
        .filter(models.Conversation.resume_id == resume_id)
        .first()
    )
    if db_conv:
        db_conv.messages = json.dumps(messages, ensure_ascii=False)
    else:
        db_conv = models.Conversation(
            resume_id=resume_id,
            messages=json.dumps(messages, ensure_ascii=False),
        )
        db.add(db_conv)
    db.commit()
    return {"message": "Conversation saved"}


@router.get("/chat/job-conversations/{job_id}")
def get_job_conversation(job_id: int, db: Session = Depends(get_db)):
    """Get conversation for a job. Returns empty messages if none saved."""
    db_conv = (
        db.query(models.JobConversation)
        .filter(models.JobConversation.job_id == job_id)
        .first()
    )
    if not db_conv:
        return {"job_id": job_id, "messages": []}
    return {"job_id": job_id, "messages": json.loads(db_conv.messages)}


@router.post("/chat/job-conversations")
def save_job_conversation(body: schemas.JobConversationSave, db: Session = Depends(get_db)):
    job_id = body.job_id
    messages = body.messages
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")
    db_conv = (
        db.query(models.JobConversation)
        .filter(models.JobConversation.job_id == job_id)
        .first()
    )
    if db_conv:
        db_conv.messages = json.dumps(messages, ensure_ascii=False)
    else:
        db_conv = models.JobConversation(
            job_id=job_id,
            messages=json.dumps(messages, ensure_ascii=False),
        )
        db.add(db_conv)
    db.commit()
    return {"message": "Job conversation saved"}


# ══════════════════════════════════════════════════════════════
# 4. export 路由
# ══════════════════════════════════════════════════════════════

@router.get("/export/markdown/{resume_id}")
def export_markdown(resume_id: int, db: Session = Depends(get_db)):
    db_resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")
    content = _clean_content(db_resume.content or "")
    buffer = io.BytesIO(content.encode("utf-8"))
    filename = _safe_filename(db_resume.title, "md", f"resume_{resume_id}.md")
    return StreamingResponse(
        buffer,
        media_type="text/markdown; charset=utf-8",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


@router.get("/export/pdf/{resume_id}")
def export_pdf(
    resume_id: int,
    db: Session = Depends(get_db),
    font_size: int = 10,
    margin_cm: float = 2.0,
):
    """font_size: 9-12, margin_cm: 1.5, 2.0, 2.5"""
    font_size = max(9, min(12, font_size))
    margin_cm = max(1.0, min(3.0, margin_cm))
    db_resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")

    content = _clean_content(db_resume.content)
    buffer = io.BytesIO()
    m = margin_cm * cm
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=m, leftMargin=m, topMargin=m, bottomMargin=m)
    doc.build(md_to_pdf_content(content, font_size))
    buffer.seek(0)

    filename = _safe_filename(db_resume.title, "pdf", f"resume_{resume_id}.pdf")
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": _content_disposition(filename)})


@router.get("/export/pdf-preview/{resume_id}")
def preview_pdf(
    resume_id: int,
    db: Session = Depends(get_db),
    font_size: int = 10,
    margin_cm: float = 2.0,
):
    response = export_pdf(resume_id, db, font_size, margin_cm)
    disposition = response.headers["Content-Disposition"]
    encoded_name = disposition.split("filename*=UTF-8''")[-1]
    response.headers["Content-Disposition"] = f"inline; filename*=UTF-8''{encoded_name}"
    return response


@router.get("/export/word/{resume_id}")
def export_word(
    resume_id: int,
    db: Session = Depends(get_db),
    font_size: int = 11,
    margin_cm: float = 2.0,
):
    """font_size: 10, 11, 12, 14 (pt). margin_cm: 1.5, 2.0, 2.5"""
    font_size = max(9, min(14, font_size))
    margin_cm = max(1.0, min(3.0, margin_cm))
    db_resume = db.query(models.Resume).filter(models.Resume.id == resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")

    content = _clean_content(db_resume.content)
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(margin_cm)
        section.bottom_margin = Cm(margin_cm)
        section.left_margin = Cm(margin_cm)
        section.right_margin = Cm(margin_cm)
    for line in content.splitlines():
        text = line.strip()
        if not text:
            doc.add_paragraph("")
            continue
        clean = _strip_md_inline(text)
        h1_pt = min(22, font_size + 8)
        h2_pt = min(16, font_size + 4)
        h3_pt = min(14, font_size + 2)
        if text.startswith("# "):
            p = doc.add_heading(clean[2:].strip(), level=1)
            for run in p.runs:
                run.font.size = Pt(h1_pt)
        elif text.startswith("## "):
            p = doc.add_heading(clean[3:].strip(), level=2)
            for run in p.runs:
                run.font.size = Pt(h2_pt)
        elif text.startswith("### "):
            p = doc.add_heading(clean[4:].strip(), level=3)
            for run in p.runs:
                run.font.size = Pt(h3_pt)
        elif text.startswith("- ") or text.startswith("* "):
            _add_paragraph_with_font(doc, clean[2:].strip(), style="List Bullet", font_pt=font_size)
        else:
            _add_paragraph_with_font(doc, clean, font_pt=font_size)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = _safe_filename(db_resume.title, "docx", f"resume_{resume_id}.docx")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


# ══════════════════════════════════════════════════════════════
# 5. settings 路由
# ══════════════════════════════════════════════════════════════

@router.get("/settings")
def get_settings():
    settings = load_settings()
    masked = {}
    custom_cfg = settings.get("custom", {})
    provider = settings.get("provider", "")

    custom_models_masked = []
    for cm in settings.get("custom_models", []):
        item = dict(cm)
        item["api_key_set"] = "••••••" if cm.get("api_key") else ""
        item.pop("api_key", None)
        custom_models_masked.append(item)

    result = {
        "provider": provider,
        "model": settings.get("model", ""),
        "api_keys_set": masked,
        "providers": PROVIDERS,
        "custom": {
            "base_url": custom_cfg.get("base_url", ""),
            "api_key_set": "••••••" if custom_cfg.get("api_key", "") else "",
            "model": custom_cfg.get("model", ""),
        },
        "custom_models": custom_models_masked,
    }

    if is_cloud_provider(provider):
        result["is_cloud"] = True
        result["cloud_template_id"] = get_cloud_template_id(provider)

    return result


@router.post("/settings/custom-models")
def create_custom_model(body: CustomModelCreate):
    """新增一个自定义模型并立即返回脱敏后的数据。"""
    settings = load_settings()
    new_cm = add_custom_model(body.model_dump(), settings)
    # 脱敏
    resp = dict(new_cm)
    resp["api_key_set"] = "••••••" if new_cm.get("api_key") else ""
    resp.pop("api_key", None)
    return resp


@router.delete("/settings/custom-models/{model_id}")
def remove_custom_model(model_id: str):
    """根据 id 删除一个自定义模型。"""
    settings = load_settings()
    deleted = delete_custom_model(model_id, settings)
    if not deleted:
        raise HTTPException(status_code=404, detail=f"Custom model not found: {model_id}")
    return {"message": "Custom model deleted"}


@router.put("/settings/custom-models/{model_id}")
def update_custom_model_route(model_id: str, body: CustomModelUpdate):
    """更新一个自定义模型。api_key 留空或 '••••••' 表示保留原值。"""
    settings = load_settings()
    updated = update_custom_model(model_id, body.model_dump(exclude_unset=True), settings)
    if not updated:
        raise HTTPException(status_code=404, detail=f"Custom model not found: {model_id}")
    # 如果当前选中的正是被编辑的模型，同步更新 settings.model（model_id 可能变了）
    if settings.get("provider") == f"{CUSTOM_PROVIDER_KEY}:{model_id}":
        settings["model"] = updated.get("model_id", "")
        save_settings(settings)
    resp = dict(updated)
    resp["api_key_set"] = "••••••" if updated.get("api_key") else ""
    resp.pop("api_key", None)
    return resp


@router.post("/settings/test-custom-model")
async def test_custom_model_direct(body: CustomModelTest):
    """
    通过后端代理发起连通性测试，避免浏览器 CORS 拦截。
    完整透传 HTTP 状态码、响应体、错误类型，方便前端展示。
    编辑模式下 api_key 为空或 '••••••' 时，用 existing_model_id 查找已存储的密钥。
    """
    # 解析 api_key：空 / 掩码 → 用已存储的密钥（需提供 existing_model_id）
    api_key = (body.api_key or "").strip()
    if (not api_key or api_key == "••••••") and body.existing_model_id:
        settings = load_settings()
        cm = get_custom_model(f"{CUSTOM_PROVIDER_KEY}:{body.existing_model_id}", settings)
        if cm:
            api_key = cm.get("api_key", "")
    if not api_key:
        return {
            "success": False,
            "error_type": "no_api_key",
            "message": "请填写 API 密钥后再测试。",
        }

    base = (body.base_url or "").strip()
    if body.full_url:
        # 用户填写完整 URL，直接使用
        test_url = base
    else:
        # 去掉末尾斜杠后拼接 /chat/completions
        clean = base.rstrip("/")
        test_url = f"{clean}/chat/completions"

    payload = json.dumps({
        "model": body.model_id,
        "messages": [{"role": "user", "content": "hi"}],
        "max_tokens": 1,
    }).encode("utf-8")

    req = urllib.request.Request(
        test_url,
        data=payload,
        method="POST",
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {api_key}",
        },
    )

    def _do_request() -> dict:
        try:
            with urllib.request.urlopen(req, timeout=15) as resp:
                body_text = resp.read().decode("utf-8", errors="replace")
                return {
                    "success": True,
                    "status_code": resp.status,
                    "body": body_text,
                }
        except urllib.error.HTTPError as e:
            # 有 HTTP 响应但状态非 2xx
            try:
                body_text = e.read().decode("utf-8", errors="replace")
            except Exception:
                body_text = "(无法读取响应体)"
            return {
                "success": False,
                "status_code": e.code,
                "status_text": e.reason or "Unknown",
                "body": body_text,
            }
        except urllib.error.URLError as e:
            # 网络层错误：超时、DNS、连接拒绝等
            if isinstance(e.reason, socket.timeout):
                return {
                    "success": False,
                    "error_type": "timeout",
                    "message": "请求超时（15秒），请检查网络或 API 地址是否正确。",
                }
            return {
                "success": False,
                "error_type": "URLError",
                "message": f"{type(e.reason).__name__}: {e.reason}",
            }
        except socket.timeout:
            return {
                "success": False,
                "error_type": "timeout",
                "message": "请求超时（15秒），请检查网络或 API 地址是否正确。",
            }
        except Exception as e:
            return {
                "success": False,
                "error_type": type(e).__name__,
                "message": str(e),
            }

    # urllib 是同步阻塞，放到线程池避免阻塞事件循环
    result = await asyncio.to_thread(_do_request)
    return result


@router.post("/settings/test")
async def test_api_connection(body: SettingsTest):
    """Test if the configured API key can connect successfully."""
    success, message = await test_connection(
        provider=body.provider,
        model=body.model,
        api_key_override=body.api_key,
        base_url_override=body.base_url,
    )
    return {"success": success, "message": message}


@router.put("/settings")
def update_settings(body: SettingsUpdate):
    settings = load_settings()
    settings["provider"] = body.provider
    settings["model"] = body.model

    if body.api_keys:
        for k, v in body.api_keys.items():
            if v == "••••••":
                continue
            settings["api_keys"][k] = v

    # 自定义模型配置
    custom_cfg = settings.get("custom", {})
    if body.custom_base_url is not None:
        custom_cfg["base_url"] = body.custom_base_url.strip()
    if body.custom_api_key is not None and body.custom_api_key != "••••••":
        custom_cfg["api_key"] = body.custom_api_key.strip()
    if body.custom_model is not None:
        custom_cfg["model"] = body.custom_model.strip()
    settings["custom"] = custom_cfg

    save_settings(settings)
    return {"message": "Settings saved"}


@router.delete("/settings/api-key/{provider}")
def clear_api_key(provider: str):
    settings = load_settings()
    if provider in settings.get("api_keys", {}):
        settings["api_keys"][provider] = ""
        save_settings(settings)
    return {"message": f"API key for {provider} cleared"}


# ══════════════════════════════════════════════════════════════
# 6. uploads 路由
# ══════════════════════════════════════════════════════════════

_AVATAR_MIME_EXT = {
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/gif": "gif",
    "image/webp": "webp",
    "image/bmp": "bmp",
}


@router.post("/uploads/avatar")
async def upload_avatar(file: UploadFile = File(...)):
    """上传候选人头像，固定保存到安装目录 data/avatars 下。

    返回本地绝对路径；生成简历时把该绝对路径发给模型，由模型在头像框中
    以 <img src="绝对路径"> 引用（预览/导出层负责渲染本地路径图片）。
    """
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="头像文件为空")
    if len(content) > 8 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="头像文件超过 8MB 限制")
    ext = _AVATAR_MIME_EXT.get((file.content_type or "").lower())
    if not ext:
        raise HTTPException(status_code=400, detail="仅支持 png / jpg / gif / webp / bmp 格式图片")

    AVATARS_DIR.mkdir(parents=True, exist_ok=True)
    avatar_file = AVATARS_DIR / f"avatar_{int(time.time() * 1000)}.{ext}"
    avatar_file.write_bytes(content)
    return {"path": str(avatar_file.resolve()), "filename": avatar_file.name}


@router.get("/avatars/{filename}")
def get_avatar_file(filename: str):
    """输出头像图片，供前端 <img> 预览（无法携带请求头，令牌走 ?token= 查询参数）。"""
    if "/" in filename or "\\" in filename or ".." in filename:
        raise HTTPException(status_code=400, detail="Invalid filename")
    p = AVATARS_DIR / filename
    if not p.is_file():
        raise HTTPException(status_code=404, detail="Avatar not found")
    return FileResponse(p)


@router.post("/uploads/parse-resume-background", response_model=ParseResumeBackgroundResponse)
async def parse_resume_background(file: UploadFile = File(...)):
    """简历 PDF → 候选人背景 Markdown。使用模型管理中当前配置的模型，扫描件需多模态模型。"""
    name = (file.filename or "").lower()
    if not name.endswith(".pdf"):
        raise HTTPException(status_code=400, detail="仅支持 PDF 简历文件")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空")

    # 内部会同步调用云端模型，放线程池避免阻塞事件循环
    text, parser, warning = await asyncio.to_thread(parse_resume_pdf_to_background, content)
    return ParseResumeBackgroundResponse(
        filename=file.filename or "resume.pdf",
        text=text,
        parser=parser,
        warning=warning,
    )


@router.post("/uploads/parse-resume-background/stream")
async def parse_resume_background_stream(file: UploadFile = File(...)):
    """流式解析简历 PDF，逐块返回文本（SSE）。使用模型管理中当前配置的模型。"""
    name = (file.filename or "").lower()
    if not name.endswith(".pdf"):
        raise HTTPException(status_code=400, detail="仅支持 PDF 简历文件")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空")

    return StreamingResponse(
        sse_with_heartbeat(stream_resume_pdf_to_background(content)),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.post("/uploads/parse-job", response_model=JobParseResponse)
async def parse_job(file: UploadFile = File(...)):
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空")
    try:
        text, _ = await asyncio.to_thread(extract_text_from_file, file.filename or "", content)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"JD解析失败: {exc}") from exc
    if not text:
        raise HTTPException(status_code=400, detail="未提取到JD文本")
    return _parse_job_fields(text)


class UploadResumeRequest(BaseModel):
    job_id: int
    accent: Optional[str] = "#2563eb"


@router.post("/uploads/upload-resume", response_model=schemas.ResumeResponse)
async def upload_resume(
    job_id: int = Query(...),
    accent: Optional[str] = Query("#2563eb"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """上传简历文件（图片/PDF/Word），自动提取文本 → LLM解析结构化JSON → LLM生成HTML简历 → 写入磁盘。

    完整链路：文件提取 → 结构化 → HTML生成，一步到位。
    """
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    content_bytes = await file.read()
    if not content_bytes:
        raise HTTPException(status_code=400, detail="上传文件为空")

    # Step 1: 提取文本（图片路径会同步调用模型，放线程池避免阻塞事件循环）
    try:
        text, parser = await asyncio.to_thread(extract_text_from_file, file.filename or "", content_bytes)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"文件解析失败: {exc}") from exc

    # docx/PDF 文本提取不足时，走 AI 多模态识别路径（与 PDF VL 路径一致）
    if not text or len(text.strip()) < 10:
        lower_name = (file.filename or "").lower()
        if lower_name.endswith((".doc", ".docx")):
            pdf_bytes = _docx_to_pdf_bytes(content_bytes)
            if pdf_bytes:
                text = _extract_from_pdf(pdf_bytes)
                parser = "docx→pdf"
                if not text or len(text.strip()) < 10:
                    text = await _ai_vl_extract_from_pdf(pdf_bytes)
                    parser = "docx→pdf+ai_vl"
        elif lower_name.endswith(".pdf"):
            text = await _ai_vl_extract_from_pdf(content_bytes)
            parser = "pdf+ai_vl"
        if not text or len(text.strip()) < 10:
            raise HTTPException(status_code=400, detail="未提取到有效文本，请检查文件清晰度")

    # Step 2: LLM 解析为结构化 JSON
    jd = _build_job_content(db_job).strip()
    parse_msg = f"请将以下简历内容提取为结构化 JSON：\n\n{text[:120000]}"
    raw_json = await complete_response(SYSTEM_PARSE_RESUME, [{"role": "user", "content": parse_msg}])
    structured_text = _strip_code_fence(raw_json).strip()

    # Step 3: LLM 根据结构化 JSON + JD 生成 HTML
    accent_val = accent if accent in ACCENTS else "#2563eb"
    gen_msg = (
        f"## 目标岗位 JD\n\n{jd}\n\n---\n\n"
        f"## 候选人结构化信息\n\n{structured_text}\n\n---\n\n"
        f"请使用主题色 {accent_val}，根据以上信息生成一份完整简历 HTML。"
    )
    gen_msg = _append_layout_directive(gen_msg, "classic", accent_val)
    raw_html = await complete_response(SYSTEM_GENERATE_RESUME, [{"role": "user", "content": gen_msg}])
    html = _strip_code_fence(raw_html).strip()

    if not html or ("<!DOCTYPE" not in html and "<html" not in html):
        raise HTTPException(status_code=502, detail="模型未生成有效的 HTML 简历，请重试或更换模型")

    # Step 4: 保存到 DB + 磁盘
    q = db.query(models.Resume).filter(models.Resume.job_id == job_id)
    angle_index = q.count() + 1
    angle = f"角度{angle_index}"
    existing_count = q.filter(models.Resume.angle == angle).count()
    title = _build_resume_title_with_angle(db_job, file.filename, angle)

    word_count = len(text)

    db_resume = models.Resume(
        job_id=job_id,
        title=title,
        content=text,
        structured=structured_text,
        rendered_html=html,
        current_template="classic",
        angle=angle,
        version=existing_count + 1,
        file_type=parser,
        file_size=len(content_bytes),
    )
    db.add(db_resume)
    db.commit()
    db.refresh(db_resume)

    write_resume_file(db_resume.id, html)

    return db_resume


@router.post("/uploads/upload-resume/stream")
async def upload_resume_stream(
    job_id: int = Query(...),
    accent: Optional[str] = Query("#2563eb"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """流式版：上传简历文件 → 提取文本 → 结构化 → 生成HTML，逐步推送进度。"""
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    content_bytes = await file.read()
    if not content_bytes:
        raise HTTPException(status_code=400, detail="上传文件为空")

    jd = _build_job_content(db_job).strip()
    accent_val = accent if accent in ACCENTS else "#2563eb"
    filename = file.filename or ""

    async def _stream():
        # Step 0: 提取文本（图片路径会同步调用模型，放线程池避免阻塞事件循环）
        try:
            text, parser = await asyncio.to_thread(extract_text_from_file, filename, content_bytes)
        except Exception as exc:
            yield f"data: {json.dumps({'type': 'error', 'content': f'文件解析失败: {exc}'})}\n\n"
            return

        # 文本提取不足时走 AI 多模态识别路径（与 PDF VL 路径一致）
        if not text or len(text.strip()) < 10:
            lower_name = filename.lower()
            if lower_name.endswith((".doc", ".docx")):
                yield f"data: {json.dumps({'type': 'progress', 'phase': 'vl', 'content': '正在转换并识别文档...'}, ensure_ascii=False)}\n\n"
                pdf_bytes = _docx_to_pdf_bytes(content_bytes)
                if pdf_bytes:
                    text = _extract_from_pdf(pdf_bytes)
                    parser = "docx→pdf"
                    if not text or len(text.strip()) < 10:
                        try:
                            text = await _ai_vl_extract_from_pdf(pdf_bytes)
                            parser = "docx→pdf+ai_vl"
                        except HTTPException as e:
                            yield f"data: {json.dumps({'type': 'error', 'content': str(e.detail)})}\n\n"
                            return
            elif lower_name.endswith(".pdf"):
                yield f"data: {json.dumps({'type': 'progress', 'phase': 'vl', 'content': '正在识别文档...'}, ensure_ascii=False)}\n\n"
                try:
                    text = await _ai_vl_extract_from_pdf(content_bytes)
                    parser = "pdf+ai_vl"
                except HTTPException as e:
                    yield f"data: {json.dumps({'type': 'error', 'content': str(e.detail)})}\n\n"
                    return
            if not text or len(text.strip()) < 10:
                yield f"data: {json.dumps({'type': 'error', 'content': '未提取到有效文本，请检查文件清晰度'})}\n\n"
                return

        # Phase 1: 结构化
        yield f"data: {json.dumps({'type': 'progress', 'phase': 'parse', 'content': '正在解析简历内容...'}, ensure_ascii=False)}\n\n"

        try:
            raw_json = await complete_response(
                SYSTEM_PARSE_RESUME,
                [{"role": "user", "content": f"请将以下简历内容提取为结构化 JSON：\n\n{text[:120000]}"}],
            )
            structured_text = _strip_code_fence(raw_json).strip()
        except Exception as e:
            yield f"data: {json.dumps({'type': 'error', 'content': f'结构化解析失败: {str(e)[:300]}'})}\n\n"
            return

        yield f"data: {json.dumps({'type': 'progress', 'phase': 'generate', 'content': '正在生成简历HTML...'}, ensure_ascii=False)}\n\n"

        # Phase 2: 生成 HTML（流式输出）
        gen_msg = (
            f"## 目标岗位 JD\n\n{jd}\n\n---\n\n"
            f"## 候选人结构化信息\n\n{structured_text}\n\n---\n\n"
            f"请使用主题色 {accent_val}，根据以上信息生成一份完整简历 HTML。"
        )
        gen_msg = _append_layout_directive(gen_msg, "classic", accent_val)
        html_parts: list[str] = []
        try:
            async for chunk in stream_response(SYSTEM_GENERATE_RESUME, [{"role": "user", "content": gen_msg}]):
                html_parts.append(chunk)
                yield f"data: {json.dumps({'type': 'html', 'content': chunk}, ensure_ascii=False)}\n\n"
        except Exception as e:
            yield f"data: {json.dumps({'type': 'error', 'content': f'HTML生成中断: {str(e)[:300]}'})}\n\n"
            return

        try:
            full_html = _strip_code_fence("".join(html_parts)).strip()
        except ValueError as exc:
            yield f"data: {json.dumps({'type': 'error', 'content': str(exc)}, ensure_ascii=False)}\n\n"
            return

        if full_html and ("<!DOCTYPE" in full_html or "<html" in full_html):
            q = db.query(models.Resume).filter(models.Resume.job_id == job_id)
            angle_index = q.count() + 1
            angle = f"角度{angle_index}"
            existing_count = q.filter(models.Resume.angle == angle).count()
            title = _build_resume_title_with_angle(db_job, filename, angle)

            db_resume = models.Resume(
                job_id=job_id,
                title=title,
                content=text,
                structured=structured_text,
                rendered_html=full_html,
                current_template="classic",
                angle=angle,
                version=existing_count + 1,
                file_type=parser,
                file_size=len(content_bytes),
            )
            db.add(db_resume)
            db.commit()
            db.refresh(db_resume)

            write_resume_file(db_resume.id, full_html)

            yield f"data: {json.dumps({'type': 'done', 'resume_id': db_resume.id}, ensure_ascii=False)}\n\n"
        else:
            yield f"data: {json.dumps({'type': 'error', 'content': '模型未生成有效的 HTML 简历，请重试或更换模型'})}\n\n"

    return StreamingResponse(
        sse_with_heartbeat(_stream()),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )
# ══════════════════════════════════════════════════════════════

@router.get("/background/profiles", response_model=List[schemas.BackgroundProfileResponse])
def list_background_profiles(db: Session = Depends(get_db)):
    """列出全部人物背景档案；若为空则自动创建一条「默认」。"""
    rows = (
        db.query(models.UserBackground)
        .order_by(models.UserBackground.id.asc())
        .all()
    )
    if not rows:
        bg = models.UserBackground(name="默认", content="")
        db.add(bg)
        db.commit()
        db.refresh(bg)
        rows = [bg]
    return rows


@router.post("/background/profiles", response_model=schemas.BackgroundProfileResponse)
def create_background_profile(
    body: schemas.BackgroundProfileCreate,
    db: Session = Depends(get_db),
):
    name = (body.name or "").strip() or "新档案"
    bg = models.UserBackground(name=name, content=body.content or "")
    db.add(bg)
    db.commit()
    db.refresh(bg)
    return bg


@router.put("/background/profiles/{profile_id}", response_model=schemas.BackgroundProfileResponse)
def update_background_profile(
    profile_id: int,
    body: schemas.BackgroundProfileUpdate,
    db: Session = Depends(get_db),
):
    bg = db.query(models.UserBackground).filter(models.UserBackground.id == profile_id).first()
    if not bg:
        raise HTTPException(status_code=404, detail="Profile not found")
    if body.name is not None:
        t = body.name.strip()
        if t:
            bg.name = t
    if body.content is not None:
        bg.content = body.content
    db.commit()
    db.refresh(bg)
    return bg


@router.delete("/background/profiles/{profile_id}")
def delete_background_profile(profile_id: int, db: Session = Depends(get_db)):
    count = db.query(models.UserBackground).count()
    if count <= 1:
        raise HTTPException(status_code=400, detail="至少保留一份背景档案")
    bg = db.query(models.UserBackground).filter(models.UserBackground.id == profile_id).first()
    if not bg:
        raise HTTPException(status_code=404, detail="Profile not found")
    db.delete(bg)
    db.commit()
    return {"ok": True, "id": profile_id}


# ══════════════════════════════════════════════════════════════
# 8. interview-sim 路由
# ══════════════════════════════════════════════════════════════

@router.get("/interview-sim/question-categories", response_model=schemas.QuestionCategoriesResponse)
async def get_question_categories():
    """可供模拟面试勾选的题目类型（与预置 JSON、生成题库分类一致）。"""
    return schemas.QuestionCategoriesResponse(categories=list(QUESTION_CATEGORIES))


@router.get("/interview-sim/questionnaire", response_model=schemas.QuestionnaireResponse)
async def get_questionnaire(
    job_id: int,
    resume_id: int,
    background_profile_id: Optional[int] = None,
    total: int = 7,
    seed: Optional[int] = None,
    categories: Optional[List[str]] = Query(
        None,
        description="可重复传参；仅从这些类别中随机抽样；不传则全库合并后抽样",
    ),
    db: Session = Depends(get_db),
):
    """从预置分类题库 + 本岗位专属题库合并后，为本场模拟抽样题单（不调用 LLM）。"""
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    extra = _db_job_questions_as_bank(db, job_id, resume_id, background_profile_id)
    t = max(3, min(12, total))
    cat_filter = [c for c in (categories or []) if c and str(c).strip()]
    items, cats = sample_questionnaire(
        _build_job_content(db_job),
        total=t,
        seed=seed,
        max_per_category=2,
        extra_questions=extra or None,
        category_filter=cat_filter if cat_filter else None,
    )
    md = questionnaire_to_markdown(items)
    return schemas.QuestionnaireResponse(
        categories_used=cats,
        items=[
            schemas.QuestionnaireItemResponse(id=q.id, category=q.category, text=q.text)
            for q in items
        ],
        questionnaire_markdown=md,
    )


@router.get("/interview-sim/bank-preview", response_model=schemas.BankPreviewResponse)
async def get_bank_preview(
    job_id: int,
    resume_id: int,
    background_profile_id: Optional[int] = None,
    db: Session = Depends(get_db),
):
    """预览：全局预置题数量/按类统计 + 本岗位专属题全文（便于管理员核对）。"""
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    preset = load_question_bank()
    by_cat = Counter(q.category for q in preset)
    by_category = {c: by_cat.get(c, 0) for c in QUESTION_CATEGORIES}

    q = db.query(models.JobInterviewQuestion).filter(models.JobInterviewQuestion.job_id == job_id)
    q = q.filter(models.JobInterviewQuestion.resume_id == resume_id)
    if background_profile_id is not None:
        q = q.filter(models.JobInterviewQuestion.background_profile_id == background_profile_id)
    rows = q.order_by(models.JobInterviewQuestion.id.asc()).all()
    return schemas.BankPreviewResponse(
        categories_allowed=list(QUESTION_CATEGORIES),
        preset=schemas.PresetBankStats(total=len(preset), by_category=by_category),
        job_questions=[
            schemas.JobQuestionPreviewRow(id=r.id, category=r.category, text=r.text) for r in rows
        ],
    )


@router.delete("/interview-sim/job-bank", response_model=schemas.JobBankDeleteResponse)
async def clear_job_interview_bank(
    job_id: int,
    resume_id: int,
    background_profile_id: Optional[int] = None,
    db: Session = Depends(get_db),
):
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")
    q = db.query(models.JobInterviewQuestion).filter(models.JobInterviewQuestion.job_id == job_id)
    q = q.filter(models.JobInterviewQuestion.resume_id == resume_id)
    if background_profile_id is not None:
        q = q.filter(models.JobInterviewQuestion.background_profile_id == background_profile_id)
    n = q.delete()
    db.commit()
    return schemas.JobBankDeleteResponse(deleted=n)


@router.get("/interview-sim/job-bank", response_model=schemas.JobInterviewBankMetaResponse)
async def get_job_interview_bank_meta(
    job_id: int,
    resume_id: int,
    background_profile_id: Optional[int] = None,
    db: Session = Depends(get_db),
):
    db_job = db.query(models.Job).filter(models.Job.id == job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")
    n = (
        db.query(models.JobInterviewQuestion)
        .filter(models.JobInterviewQuestion.job_id == job_id)
        .filter(models.JobInterviewQuestion.resume_id == resume_id)
        .filter(
            models.JobInterviewQuestion.background_profile_id == background_profile_id
            if background_profile_id is not None
            else models.JobInterviewQuestion.background_profile_id.is_(None)
        )
        .count()
    )
    return schemas.JobInterviewBankMetaResponse(count=n)


@router.post("/interview-sim/generate-bank", response_model=schemas.GenerateInterviewBankResponse)
async def generate_job_interview_bank(
    request: schemas.GenerateInterviewBankRequest,
    db: Session = Depends(get_db),
):
    """根据 JD + 简历 + 可选背景 调用 LLM 生成专属面试题并入库；之后「开始面试」抽样时会与全局题库合并。"""
    db_job = db.query(models.Job).filter(models.Job.id == request.job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    if request.replace:
        db.query(models.JobInterviewQuestion).filter(models.JobInterviewQuestion.job_id == request.job_id).filter(
            models.JobInterviewQuestion.resume_id == request.resume_id
        ).filter(
            models.JobInterviewQuestion.background_profile_id == request.background_profile_id
        ).delete()
        db.commit()

    try:
        db_resume = db.query(models.Resume).filter(models.Resume.id == request.resume_id).first()
        if not db_resume:
            raise HTTPException(status_code=404, detail="Resume not found")
        db_bg = None
        if request.background_profile_id is not None:
            db_bg = db.query(models.UserBackground).filter(models.UserBackground.id == request.background_profile_id).first()
        dicts = await generate_job_question_dicts(
            _build_job_content(db_job),
            db_resume.content or "",
            db_bg.content if db_bg else None,
        )
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"LLM 生成失败: {str(e)[:200]}")

    if not dicts:
        raise HTTPException(status_code=502, detail="未生成有效题目，请检查模型配置后重试")

    for d in dicts:
        db.add(
            models.JobInterviewQuestion(
                job_id=request.job_id,
                resume_id=request.resume_id,
                background_profile_id=request.background_profile_id,
                category=d["category"],
                text=d["text"],
            )
        )
    db.commit()

    total = (
        db.query(models.JobInterviewQuestion)
        .filter(models.JobInterviewQuestion.job_id == request.job_id)
        .filter(models.JobInterviewQuestion.resume_id == request.resume_id)
        .filter(models.JobInterviewQuestion.background_profile_id == request.background_profile_id)
        .count()
    )
    return schemas.GenerateInterviewBankResponse(added=len(dicts), total_for_job=total)


@router.put("/interview-sim/job-question", response_model=schemas.JobQuestionUpdateResponse)
async def update_job_question(
    request: schemas.JobQuestionUpdateRequest,
    db: Session = Depends(get_db),
):
    q = (
        db.query(models.JobInterviewQuestion)
        .filter(models.JobInterviewQuestion.id == request.question_id)
        .filter(models.JobInterviewQuestion.job_id == request.job_id)
        .filter(models.JobInterviewQuestion.resume_id == request.resume_id)
        .filter(models.JobInterviewQuestion.background_profile_id == request.background_profile_id)
        .first()
    )
    if not q:
        raise HTTPException(status_code=404, detail="Question not found")
    q.text = request.text
    db.commit()
    return schemas.JobQuestionUpdateResponse(updated=1)


@router.delete("/interview-sim/job-question", response_model=schemas.JobQuestionDeleteResponse)
async def delete_job_question(
    request: schemas.JobQuestionDeleteRequest,
    db: Session = Depends(get_db),
):
    n = (
        db.query(models.JobInterviewQuestion)
        .filter(models.JobInterviewQuestion.id == request.question_id)
        .filter(models.JobInterviewQuestion.job_id == request.job_id)
        .filter(models.JobInterviewQuestion.resume_id == request.resume_id)
        .filter(models.JobInterviewQuestion.background_profile_id == request.background_profile_id)
        .delete()
    )
    db.commit()
    return schemas.JobQuestionDeleteResponse(deleted=n)


@router.post("/interview-sim/stream")
async def interview_sim_stream(request: schemas.InterviewSimRequest, db: Session = Depends(get_db)):
    db_job = db.query(models.Job).filter(models.Job.id == request.job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    db_resume = db.query(models.Resume).filter(models.Resume.id == request.resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")
    if db_resume.job_id != request.job_id:
        raise HTTPException(status_code=400, detail="Resume does not belong to this job")

    messages = [msg.model_dump() for msg in request.messages]

    return StreamingResponse(
        sse_with_heartbeat(_stream_sim(
            job_content=_build_job_content(db_job),
            resume_content=db_resume.content or "",
            messages=messages,
            user_background=request.user_background,
            questionnaire_markdown=request.questionnaire_markdown,
            job_id=request.job_id,
            resume_id=request.resume_id,
        )),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.post("/interview-sim/report", response_model=schemas.InterviewReportResponse)
async def interview_sim_report(request: schemas.InterviewReportRequest, db: Session = Depends(get_db)):
    db_job = db.query(models.Job).filter(models.Job.id == request.job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")

    db_resume = db.query(models.Resume).filter(models.Resume.id == request.resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")
    if db_resume.job_id != request.job_id:
        raise HTTPException(status_code=400, detail="Resume does not belong to this job")

    job_content = _build_job_content(db_job)
    resume_content = db_resume.content or ""
    transcript = _format_transcript([m.model_dump() for m in request.messages])
    if not transcript.strip():
        raise HTTPException(status_code=400, detail="No interview content to analyze")

    context_parts = [
        f"## 岗位 JD\n\n{job_content}",
        f"## 候选人简历\n\n{resume_content}",
        f"## 模拟面试对话记录\n\n{transcript}",
    ]
    if request.user_background:
        context_parts.append(f"## 候选人补充经历\n\n{request.user_background}")

    context = "\n\n---\n\n".join(context_parts)
    system = f"{INTERVIEW_REPORT_SYSTEM}\n\n---\n\n{context}"

    user_msg = "请根据以上材料生成面试复盘报告（Markdown）。直接输出报告正文，不要前言套话。"
    report_md = await complete_response(system, [{"role": "user", "content": user_msg}], source="interview")

    return schemas.InterviewReportResponse(report=report_md)


# ══════════════════════════════════════════════════════════════
# 9. evaluation 路由
# ══════════════════════════════════════════════════════════════

@router.post("/evaluation/scorecard", response_model=schemas.EvaluationScorecardResponse)
async def generate_scorecard(request: schemas.EvaluationScorecardRequest, db: Session = Depends(get_db)):
    db_job = db.query(models.Job).filter(models.Job.id == request.job_id).first()
    if not db_job:
        raise HTTPException(status_code=404, detail="Job not found")
    db_resume = db.query(models.Resume).filter(models.Resume.id == request.resume_id).first()
    if not db_resume:
        raise HTTPException(status_code=404, detail="Resume not found")

    parts = [
        f"## 岗位 JD\n\n{_build_job_content(db_job)}",
        f"## 候选人简历\n\n{_extract_resume_text(db_resume)}",
    ]
    if request.user_background:
        parts.append(f"## 候选人背景补充\n\n{request.user_background}")
    if request.transcript:
        parts.append(f"## 模拟面试对话摘录\n\n{request.transcript}")
    competency_set = _pick_competency_set_by_profile(getattr(db_job, "competency_profile", "default"), db_job.title or "")
    competency_text = "、".join(competency_set) if competency_set else "综合能力、岗位匹配、风险识别、改进建议"
    user_msg = (
        "\n\n---\n\n".join(parts)
        + f"\n\n请优先围绕以下能力项生成评分卡：{competency_text}\n\n请输出评分卡 JSON。"
    )

    try:
        raw = await complete_response(EVAL_SYSTEM_PROMPT, [{"role": "user", "content": user_msg}])
        data = json.loads(_strip_code_fence(raw, json_mode=True))
        if not isinstance(data, dict):
            raise ValueError("invalid json object")
        normalized = _normalize_scorecard(data)
    except Exception as e:
        import logging
        logging.getLogger(__name__).warning("评分卡生成失败 raw=%s", raw[:500] if 'raw' in dir() else 'N/A', exc_info=True)
        raise HTTPException(status_code=502, detail=f"评分卡生成失败: {str(e)[:180]}")

    db.add(
        models.EvaluationReport(
            job_id=request.job_id,
            resume_id=request.resume_id,
            report_type="scorecard",
            content_json=json.dumps(normalized.model_dump(), ensure_ascii=False),
        )
    )
    db.commit()
    return normalized


@router.get("/evaluation/scorecard-history", response_model=List[schemas.EvaluationReportListItem])
def list_scorecard_history(job_id: int, resume_id: Optional[int] = None, limit: int = 20, db: Session = Depends(get_db)):
    limit = max(1, min(100, int(limit)))
    q = (
        db.query(models.EvaluationReport)
        .filter(models.EvaluationReport.job_id == job_id)
        .filter(models.EvaluationReport.report_type == "scorecard")
    )
    if resume_id is not None:
        q = q.filter(models.EvaluationReport.resume_id == resume_id)
    return q.order_by(models.EvaluationReport.created_at.desc()).limit(limit).all()


@router.get("/evaluation/scorecard-history/{report_id}", response_model=schemas.EvaluationReportListItem)
def get_scorecard_history_item(report_id: int, db: Session = Depends(get_db)):
    row = db.query(models.EvaluationReport).filter(models.EvaluationReport.id == report_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Scorecard report not found")
    return row


# ══════════════════════════════════════════════════════════════
# AI 话术生成（/agent/greeting-generator）
# ══════════════════════════════════════════════════════════════

class GreetingStyleOption(BaseModel):
    id: str
    name: str
    description: str

GREETING_STYLES = [
    GreetingStyleOption(id="experience-match", name="经验匹配型", description="直接亮年限+核心技能，对标JD要求"),
    GreetingStyleOption(id="project-result", name="项目成果型", description="用具体项目/数据说话，突出战绩"),
    GreetingStyleOption(id="enthusiasm", name="热情态度型", description="表达对公司/岗位的强烈兴趣和认同"),
    GreetingStyleOption(id="concise-direct", name="简洁直接型", description="一句话说清我是谁+能做什么"),
    GreetingStyleOption(id="skill-keyword", name="技能关键词型", description="堆砌JD技术栈关键词，过筛选"),
    GreetingStyleOption(id="learning-potential", name="学习潜力型", description="强调学习能力和适应速度"),
]


@router.get("/greeting-generator/styles", response_model=List[GreetingStyleOption])
def list_greeting_styles():
    return GREETING_STYLES


class ParseJdImageRequest(BaseModel):
    pass


class JdParseResult(BaseModel):
    title: str = ""
    company: Optional[str] = None
    responsibilities: List[str] = []
    requirements: List[str] = []
    raw_text: str = ""


@router.post("/greeting-generator/parse-jd-image", response_model=JdParseResult)
async def parse_jd_image(file: UploadFile = File(...)):
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空")

    try:
        text, _ = await asyncio.to_thread(extract_text_from_file, file.filename or "", content)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"JD图片识别失败: {exc}") from exc

    if not text or not text.strip():
        raise HTTPException(status_code=400, detail="未识别到JD文字，请检查图片清晰度")

    parse_prompt = f"""请从以下岗位JD文本中提取结构化信息，输出严格 JSON 对象（不要 markdown 代码块、不要注释、不要 JSON 之外的任何文字）：
{{
  "title": "岗位名称",
  "company": "公司名（可选）",
  "responsibilities": ["职责1", "职责2"],
  "requirements": ["要求1", "要求2"]
}}

JD文本：
{text[:3000]}

规则：
- 提取失败的字段赋空字符串或空数组
- 职责和要求尽量拆分为独立条目
- 保留关键技术词和量化数据"""

    try:
        raw = await complete_response(
            "你是JD信息提取助手，只输出JSON，不输出其他内容。",
            [{"role": "user", "content": parse_prompt}],
        )
        data = json.loads(_strip_code_fence(raw, json_mode=True))
    except Exception:
        data = {}

    return JdParseResult(
        title=data.get("title", ""),
        company=data.get("company"),
        responsibilities=data.get("responsibilities", []),
        requirements=data.get("requirements", []),
        raw_text=text.strip(),
    )


class GenerateGreetingsRequest(BaseModel):
    jd_text: str
    style: str = "professional"
    count: int = Field(default=1, ge=1, le=8)
    existing_count: int = Field(default=0, ge=0, description="已生成的话术数量，用于区分新话术角度")
    resume: Optional[str] = None
    expect_jd: Optional[str] = None


class GreetingItem(BaseModel):
    index: int
    text: str
    style_label: str


class GenerateGreetingsResponse(BaseModel):
    greetings: List[GreetingItem]
    suggestions: List[dict]


_SYSTEM_PROMPT_GREETING = """你是一位资深求职打招呼助手。你的任务是根据招聘JD和求职者背景，生成个性化、有针对性的打招呼话术。

## 工作流程
1. 仔细分析JD，提取核心要求：岗位实际职责、必备技能栈、经验要求、行业方向
2. 将JD的核心职责和技能要求与求职者背景进行匹配
3. 根据指定风格生成打招呼话术

## 话术要求
- 每条不超过100字
- 针对JD中的具体技术要求，提及匹配的技能或经验
- 语气自然、真诚、专业，避免模板化
- 提及JD中的关键技术词或业务方向，体现契合度
- 不要出现"我看了贵公司的招聘"等套话
- 如果已有之前生成的话术，确保新话术从不同角度切入，避免与已有话术雷同

## 严禁编造（最高优先级）
- 话术中提到的技能、经验、项目、学校、工作年限等信息，必须是求职者背景中明确包含的
- 禁止根据JD反向编造求职者不具备的经验或技能
- 如果背景信息不足，宁可只写泛化表达（如"对该领域有浓厚兴趣"），也绝不凭空捏造具体经历

## 风格说明
根据用户指定的风格生成话术：
- experience-match：直接亮年限+核心技能，对标JD要求，适用于经验高度匹配的岗位
- project-result：用具体项目/数据说话，突出战绩，适用于有亮眼项目经历、想差异化竞争
- enthusiasm：表达对公司/岗位的强烈兴趣和认同，适用于经验略欠但意愿强、公司文化导向
- concise-direct：一句话说清"我是谁+能做什么"，适用于HR忙碌、批量投递时提高回复率
- skill-keyword：堆砌JD里的技术栈关键词，过筛选，适用于大厂/系统初筛、关键词匹配场景
- learning-potential：强调学习能力和适应速度，弱化经验不足，适用于转行、应届生、经验不够但潜力大

## 输出格式
严格输出 JSON 对象（不要 markdown 代码块、不要注释、不要 JSON 之外的任何文字）：
{{"text": "话术内容", "style_label": "风格标签（2-4字）"}}

style_label 体现该话术的侧重角度（如"专业型"、"成果导向"、"技能匹配"等）。"""


@router.post("/greeting-generator/generate", response_model=GenerateGreetingsResponse)
async def generate_greetings(request: GenerateGreetingsRequest):
    jd = (request.jd_text or "").strip()[:3000]
    if not jd:
        raise HTTPException(status_code=400, detail="JD内容不能为空")

    style_info = next((s for s in GREETING_STYLES if s.id == request.style), GREETING_STYLES[0])

    sections = [f"## 岗位JD信息\n{jd}"]
    if request.resume and request.resume.strip():
        sections.append(f"## 求职者真实背景（话术必须基于此信息，禁止编造未提及的经历/技能/项目）\n{request.resume.strip()[:1000]}")
    if request.expect_jd and request.expect_jd.strip():
        sections.append(f"## 期望岗位描述\n{request.expect_jd.strip()[:1000]}")
    sections.append(f"## 生成要求\n风格：{style_info.name}（{style_info.description}）")
    if request.existing_count > 0:
        sections.append(f"注意：用户已生成过 {request.existing_count} 条话术，请从新的角度切入，确保与之前的话术不重复。")
    sections.append(f"生成 {request.count} 条话术，每条从不同侧重点切入")
    user_content = "\n\n".join(sections)

    try:
        raw = await complete_response(
            _SYSTEM_PROMPT_GREETING,
            [{"role": "user", "content": user_content}],
            source="copy",
        )
        data = json.loads(_strip_code_fence(raw, json_mode=True))
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"话术生成失败: {str(e)[:200]}")

    greetings: List[GreetingItem] = []

    if isinstance(data, dict):
        text = data.get("text", "").strip()
        label = data.get("style_label", style_info.name)
        if text:
            for prefix in ("MATCH", "Match", "match", "匹配", "招呼语：", "招呼语:", "打招呼：", "打招呼:"):
                if text.startswith(prefix):
                    text = text[len(prefix):].strip()
            if len(text) >= 2 and text[0] in '"「『"' and text[-1] in '"」』"':
                text = text[1:-1].strip()
            greetings.append(GreetingItem(index=request.existing_count, text=text, style_label=label))
    elif isinstance(data, list):
        for i, item in enumerate(data[:request.count]):
            if isinstance(item, dict):
                text = item.get("text", "").strip()
                label = item.get("style_label", style_info.name)
            elif isinstance(item, str):
                text = item.strip()
                label = style_info.name
            else:
                continue
            if not text:
                continue
            for prefix in ("MATCH", "Match", "match", "匹配", "招呼语：", "招呼语:", "打招呼：", "打招呼:"):
                if text.startswith(prefix):
                    text = text[len(prefix):].strip()
            if len(text) >= 2 and text[0] in '"「『"' and text[-1] in '"」』"':
                text = text[1:-1].strip()
            greetings.append(GreetingItem(index=request.existing_count + i, text=text, style_label=label))

    if not greetings:
        raise HTTPException(status_code=502, detail="模型未生成有效话术，请重试")

    suggestions = [
        {"icon": "briefcase", "title": "突出项目经验", "description": "结合JD职责，强调相关项目成果"},
        {"icon": "target", "title": "强调匹配技能", "description": "精准对应JD要求的技能栈"},
        {"icon": "trending-up", "title": "展示学习能力", "description": "体现快速适应和持续成长"},
    ]

    return GenerateGreetingsResponse(greetings=greetings, suggestions=suggestions)


# ══════════════════════════════════════════════════════════════
# 职位分析（/agent/job-analysis）
# ══════════════════════════════════════════════════════════════

class JobAnalysisRequest(BaseModel):
    jd_text: str


class SkillItem(BaseModel):
    name: str
    level: str = "required"
    description: str = ""


class KnowledgeItem(BaseModel):
    topic: str
    description: str = ""
    resources: List[str] = []


class PlanPhase(BaseModel):
    phase: str
    duration: str
    goals: List[str] = []
    actions: List[str] = []


class InterviewQuestion(BaseModel):
    category: str
    question: str
    hints: List[str] = []


class JobAnalysisResult(BaseModel):
    job_title: str = ""
    summary: str = ""
    skills: List[SkillItem] = []
    knowledge: List[KnowledgeItem] = []
    plan: List[PlanPhase] = []
    interview_questions: List[InterviewQuestion] = []


_SYSTEM_PROMPT_JOB_ANALYSIS = """你是一位资深技术招聘顾问和职业规划师。根据用户提供的岗位JD，进行深度分析。

## 输出要求

严格输出 JSON 对象（不要 markdown 代码块、不要注释、不要 JSON 之外的任何文字），结构如下：

{
  "job_title": "岗位名称",
  "summary": "2-3句话的岗位核心概述，包括方向、难度、市场情况",
  "skills": [
    {"name": "技能名", "level": "required|preferred|bonus", "description": "为什么需要、掌握到什么程度"}
  ],
  "knowledge": [
    {"topic": "知识点", "description": "需要掌握的核心内容", "resources": ["推荐学习资源1", "推荐学习资源2"]}
  ],
  "plan": [
    {"phase": "阶段名", "duration": "预计时间", "goals": ["目标1"], "actions": ["具体行动1"]}
  ],
  "interview_questions": [
    {"category": "分类", "question": "面试问题", "hints": ["答题要点1", "答题要点2"]}
  ]
}

## 分析维度

### skills - 技能要求
从JD中提取所有技术技能和软技能，分为三级：
- required：JD明确要求的必备技能
- preferred：JD提到"优先""加分""熟悉优先"的技能
- bonus：JD未提及但该岗位通常需要的相关技能

### knowledge - 知识点学习
针对每个核心技能，拆解出需要学习的具体知识点：
- topic：知识点名称
- description：需要掌握的核心内容、深度要求
- resources：推荐的学习资源（书籍、官方文档、在线课程等，尽量具体）

### plan - 学习方案
生成一份分阶段的学习计划（3-5个阶段），从基础到进阶：
- phase：阶段名称
- duration：建议时长
- goals：该阶段要达成的目标
- actions：具体的行动步骤（学什么、做什么练习、完成什么项目）

### interview_questions - 面试问题
针对该岗位的高频面试题（8-12题），覆盖：
- 基础知识题
- 项目经验题
- 场景设计题
- 行为面试题
每题附带答题要点提示。

## 规则
1. 技能分析要全面，不要遗漏JD中隐含的技能要求
2. 学习方案要切实可行，有明确的阶段和可执行的行动
3. 面试问题要贴合该岗位实际情况，不要泛泛而谈
4. level 只能是 required / preferred / bonus 三选一
5. 中文输出"""


@router.post("/job-analysis/analyze", response_model=JobAnalysisResult)
async def analyze_job(request: JobAnalysisRequest):
    jd = (request.jd_text or "").strip()[:4000]
    if not jd:
        raise HTTPException(status_code=400, detail="JD内容不能为空")

    user_content = f"## 岗位JD信息\n\n{jd}"

    try:
        raw = await complete_response(
            _SYSTEM_PROMPT_JOB_ANALYSIS,
            [{"role": "user", "content": user_content}],
        )
        data = json.loads(_strip_code_fence(raw, json_mode=True))
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"职位分析失败: {str(e)[:200]}")

    skills = []
    for s in data.get("skills", []):
        if isinstance(s, dict) and s.get("name"):
            skills.append(SkillItem(
                name=s["name"],
                level=s.get("level", "required"),
                description=s.get("description", ""),
            ))

    knowledge = []
    for k in data.get("knowledge", []):
        if isinstance(k, dict) and k.get("topic"):
            knowledge.append(KnowledgeItem(
                topic=k["topic"],
                description=k.get("description", ""),
                resources=k.get("resources", []),
            ))

    plan = []
    for p in data.get("plan", []):
        if isinstance(p, dict) and p.get("phase"):
            plan.append(PlanPhase(
                phase=p["phase"],
                duration=p.get("duration", ""),
                goals=p.get("goals", []),
                actions=p.get("actions", []),
            ))

    questions = []
    for q in data.get("interview_questions", []):
        if isinstance(q, dict) and q.get("question"):
            questions.append(InterviewQuestion(
                category=q.get("category", ""),
                question=q["question"],
                hints=q.get("hints", []),
            ))

    return JobAnalysisResult(
        job_title=data.get("job_title", ""),
        summary=data.get("summary", ""),
        skills=skills,
        knowledge=knowledge,
        plan=plan,
        interview_questions=questions,
    )
